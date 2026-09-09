#!/usr/bin/env python3
"""
Universal TEA Engine V0.6
==========================

Workbook-driven techno-economic analysis engine for the Universal TEA V0.6
input workbooks.

Works with:
- Universal_TEA_Input_Workbook_V0_6_BASE.xlsx
- Universal_TEA_Input_Workbook_V0_6_CASE1_H2A_Biomass.xlsx
- Universal_TEA_Input_Workbook_V0_6_CASE2_PEM.xlsx
- other workbooks following the same sheet/header structure.

Outputs automatically:
- <case>_TEA_Results_V0_6.xlsx
- <case>_TEA_Report_V0_6.docx
- <case>_CAPEX.png
- <case>_OPEX.png
- <case>_Cash_Flow.png
- <case>_Sensitivity.png
- <case>_Scenarios.png
- <case>_Validation.png (when published validation targets are active)

Important V0.6 scope:
- Simple and H2A-style detailed all-equity DCF are implemented.
- Debt-financing inputs are read and reported, but NOT applied in V0.6.
- Equipment-degradation inputs are read and reported, but NOT applied in V0.6.
  Those are intentionally reserved for the next methodology upgrade.
"""

from __future__ import annotations

import argparse
import copy
import math
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Tuple, Optional

try:
    import numpy as np
    import pandas as pd
    import matplotlib.pyplot as plt
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
except ImportError as exc:
    raise SystemExit(
        "Missing Python package. In Google Colab run:\n"
        "!pip -q install openpyxl pandas numpy matplotlib python-docx\n\n"
        f"Original import error: {exc}"
    )


# ---------------------------------------------------------------------
# Generic helpers
# ---------------------------------------------------------------------

def yes(value) -> bool:
    return str(value).strip().lower() in {"yes", "y", "true", "1", "active"}

def fnum(value, default=0.0) -> float:
    if value is None or value == "":
        return float(default)
    try:
        return float(value)
    except (TypeError, ValueError):
        return float(default)

def slugify(text: str) -> str:
    text = re.sub(r"[^A-Za-z0-9]+", "_", str(text)).strip("_")
    return text[:70] or "TEA_Case"

def pct_dev(calc, ref):
    if ref in (None, "", 0):
        return np.nan
    return (float(calc) - float(ref)) / float(ref) * 100.0

def safe_div(a, b, default=np.nan):
    return a / b if b not in (0, None) else default

def npv(cashflows: List[float], rate: float) -> float:
    return float(sum(cf / ((1.0 + rate) ** t) for t, cf in enumerate(cashflows)))

def solve_irr(cashflows: List[float]) -> float:
    # Robust bisection over a wide economically meaningful interval.
    lo, hi = -0.99, 5.0
    flo, fhi = npv(cashflows, lo), npv(cashflows, hi)
    if flo == 0:
        return lo
    if fhi == 0:
        return hi
    if flo * fhi > 0:
        return np.nan
    for _ in range(120):
        mid = (lo + hi) / 2.0
        fm = npv(cashflows, mid)
        if abs(fm) < 1e-7:
            return mid
        if flo * fm <= 0:
            hi = mid
            fhi = fm
        else:
            lo = mid
            flo = fm
    return (lo + hi) / 2.0

def solve_root_price(builder, discount_rate, low=0.0, high=10.0) -> float:
    while npv(builder(high), discount_rate) < 0:
        high *= 2.0
        if high > 1e6:
            return np.nan
    for _ in range(100):
        mid = (low + high) / 2.0
        if npv(builder(mid), discount_rate) > 0:
            high = mid
        else:
            low = mid
    return (low + high) / 2.0

def macrs_20_half_year_rates() -> List[float]:
    """20-year MACRS: 150% declining balance, half-year convention."""
    basis = 1.0
    rates = []
    db_rate = 1.5 / 20.0
    for year in range(1, 22):
        period = 0.5 if year in (1, 21) else 1.0
        remaining = sum(0.5 if y in (1, 21) else 1.0 for y in range(year, 22))
        db_ded = basis * db_rate * period
        sl_ded = basis / remaining * period
        deduction = min(max(db_ded, sl_ded), basis)
        rates.append(deduction)
        basis -= deduction
    return rates

MACRS20 = macrs_20_half_year_rates()


# ---------------------------------------------------------------------
# Excel reader
# ---------------------------------------------------------------------

class WorkbookReader:
    def __init__(self, path: str):
        self.path = Path(path)
        self.wb = load_workbook(path, data_only=False, read_only=False)

    def has_sheet(self, name: str) -> bool:
        return name in self.wb.sheetnames

    def rows_by_header(self, sheet_name: str, header_row: int = 4) -> List[Dict]:
        if not self.has_sheet(sheet_name):
            return []
        ws = self.wb[sheet_name]
        headers = [ws.cell(header_row, c).value for c in range(1, ws.max_column + 1)]
        rows = []
        for r in range(header_row + 1, ws.max_row + 1):
            values = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            if all(v in (None, "") for v in values):
                continue
            row = {}
            for h, v in zip(headers, values):
                if h not in (None, ""):
                    row[str(h).strip()] = v
            rows.append(row)
        return rows

    def parameter_map(self, sheet_name: str, header_row=4,
                      key_col="Parameter", value_col="Value") -> Dict[str, object]:
        rows = self.rows_by_header(sheet_name, header_row)
        out = {}
        for row in rows:
            k = row.get(key_col)
            if k not in (None, ""):
                out[str(k).strip()] = row.get(value_col)
        return out


# ---------------------------------------------------------------------
# Model context
# ---------------------------------------------------------------------

@dataclass
class ModelContext:
    case_name: str
    process_category: str
    primary_product: str
    hours: float
    life: int
    currency: str
    reference_year: int

    product_rate: float
    product_unit: str
    selling_price: float
    selling_price_unit: str

    equipment_df: pd.DataFrame
    material_df: pd.DataFrame
    utility_df: pd.DataFrame
    waste_df: pd.DataFrame
    fixed_rows_df: pd.DataFrame

    purchased_equipment: float
    installed_equipment: float
    fci: float
    tci: float
    land: float

    material_opex: float
    electricity_opex: float
    other_utility_opex: float
    waste_opex: float
    direct_variable_opex: float
    fixed_opex: float

    target_return: float
    tax_rate: float
    depreciation_method: str
    depreciation_period: int
    wc_factor_simple: float
    startup_factor_simple: float
    salvage_factor_simple: float

    finance_mode: str
    detailed: Dict[str, float]
    warnings: List[str]

    # Information-only reference annual costs. These may be inactive in base OPEX
    # to avoid double counting, but are useful for sensitivity decomposition.
    reference_material_costs: Dict[str, float]
    reference_utility_costs: Dict[str, float]
    reference_equipment_installed: Dict[str, float]


# ---------------------------------------------------------------------
# Unit and annual-cost handling
# ---------------------------------------------------------------------

def annual_quantity(flow: float, flow_unit: str, hours: float) -> Tuple[float, str]:
    u = str(flow_unit or "").strip().lower().replace("³", "3")
    if u == "kg/h":
        return flow * hours, "kg"
    if u == "ton/h":
        return flow * hours, "ton"
    if u == "kg/year":
        return flow, "kg"
    if u == "ton/year":
        return flow, "ton"
    if u in {"nm3/h", "nm³/h"}:
        return flow * hours, "Nm3"
    if u == "m3/h":
        return flow * hours, "m3"
    if u == "m3/year":
        return flow, "m3"
    if u == "kwh/h":
        return flow * hours, "kWh"
    if u == "mwh/h":
        return flow * hours, "MWh"
    if u == "mj/h":
        return flow * hours, "MJ"
    if u == "gj/h":
        return flow * hours, "GJ"
    if u == "l/h":
        return flow * hours, "L"
    if u in {"per year", "/year", "year"}:
        return flow, "per year"
    return flow * hours, u or "unit"

def convert_quantity_for_price(qty: float, qty_unit: str, price_unit: str) -> float:
    pu = str(price_unit or "").strip().lower()
    qu = str(qty_unit or "").strip().lower()
    if "/year" in pu:
        return 1.0
    if "/kg" in pu:
        if qu == "kg":
            return qty
        if qu == "ton":
            return qty * 1000.0
    if "/ton" in pu:
        if qu == "ton":
            return qty
        if qu == "kg":
            return qty / 1000.0
    if "/nm3" in pu or "/nm³" in pu:
        return qty
    if "/m3" in pu:
        return qty
    if "/kwh" in pu:
        if qu == "kwh":
            return qty
        if qu == "mwh":
            return qty * 1000.0
    if "/mwh" in pu:
        if qu == "mwh":
            return qty
        if qu == "kwh":
            return qty / 1000.0
    if "/gj" in pu:
        if qu == "gj":
            return qty
        if qu == "mj":
            return qty / 1000.0
    return qty

def annual_cost(flow: float, flow_unit: str, price: float, price_unit: str, hours: float) -> float:
    qty, q_unit = annual_quantity(flow, flow_unit, hours)
    return convert_quantity_for_price(qty, q_unit, price_unit) * price


# ---------------------------------------------------------------------
# Parse workbook into model context
# ---------------------------------------------------------------------

def parse_model(path: str) -> Tuple[ModelContext, WorkbookReader, pd.DataFrame, pd.DataFrame]:
    rd = WorkbookReader(path)
    project = rd.parameter_map("Project")
    financial = rd.rows_by_header("Financial")
    detailed_rows = rd.rows_by_header("Detailed_Finance")
    validation_rows = rd.rows_by_header("Validation_Targets")

    case_name = str(project.get("Project name") or Path(path).stem)
    process_category = str(project.get("Process category") or "Unspecified")
    primary_product = str(project.get("Primary product") or "Primary product")
    hours = fnum(project.get("Annual operating hours"), 8000)
    life = int(round(fnum(project.get("Project operating life"), 20)))
    currency = str(project.get("Currency") or "USD")
    reference_year = int(round(fnum(project.get("Currency reference year"), 2025)))

    # Products
    products = rd.rows_by_header("Products")
    primary_row = None
    for row in products:
        if yes(row.get("Active?")) and str(row.get("Role", "")).strip().lower() == "primary":
            primary_row = row
            break
    if primary_row is None:
        for row in products:
            if yes(row.get("Active?")):
                primary_row = row
                break
    if primary_row is None:
        raise ValueError("No active product row found in Products sheet.")

    product_rate = fnum(primary_row.get("Production Rate"))
    product_unit = str(primary_row.get("Unit") or "")
    selling_price = fnum(primary_row.get("Selling Price / Credit"))
    selling_price_unit = str(primary_row.get("Price Unit") or "")

    # Equipment
    equipment_rows = rd.rows_by_header("Equipment")
    eq_records = []
    purchased_total = 0.0
    installed_total = 0.0
    land = 0.0
    reference_equipment_installed = {}
    for row in equipment_rows:
        if not yes(row.get("Active?")):
            continue
        name = str(row.get("Equipment / Cost Block") or "Equipment")
        actual = fnum(row.get("Actual Capacity"), 1.0)
        base_cap = fnum(row.get("Base Capacity"), 1.0)
        base_cost = fnum(row.get("Base Cost"), 0.0)
        exponent = fnum(row.get("Scaling Exponent"), 1.0)
        install_factor = fnum(row.get("Installation Factor"), 1.0)
        if base_cap <= 0:
            base_cap = 1.0
        purchased = base_cost * (actual / base_cap) ** exponent
        installed = purchased * install_factor
        purchased_total += purchased
        installed_total += installed
        if name.strip().lower() == "land":
            land += installed
        reference_equipment_installed[name.lower()] = installed
        eq_records.append({
            "Equipment": name,
            "Actual Capacity": actual,
            "Capacity Unit": row.get("Capacity Unit"),
            "Base Capacity": base_cap,
            "Base Cost": base_cost,
            "Scaling Exponent": exponent,
            "Installation Factor": install_factor,
            "Purchased Cost": purchased,
            "Installed Cost": installed,
            "Source ID": row.get("Source ID"),
        })
    equipment_df = pd.DataFrame(eq_records)

    # Financial/simple layer
    fparam = {}
    factive = {}
    for row in financial:
        p = str(row.get("Parameter") or "").strip()
        if p:
            fparam[p] = row.get("Value")
            factive[p] = yes(row.get("Active?"))

    target_return = fnum(fparam.get("Discount rate / target return"), 0.08)
    tax_rate = fnum(fparam.get("Corporate tax rate"), 0.25)
    depreciation_method = str(fparam.get("Depreciation method") or "Straight-line")
    depreciation_period = int(round(fnum(fparam.get("Depreciation period"), 10)))
    wc_factor_simple = fnum(fparam.get("Working capital factor"), 0.0) if factive.get("Working capital factor", True) else 0.0
    startup_factor_simple = fnum(fparam.get("Startup capital factor"), 0.0) if factive.get("Startup capital factor", True) else 0.0
    salvage_factor_simple = fnum(fparam.get("Salvage factor"), 0.0) if factive.get("Salvage factor", True) else 0.0
    eng_factor = fnum(fparam.get("Engineering / indirect factor"), 0.0) if factive.get("Engineering / indirect factor", True) else 0.0
    proj_cont = fnum(fparam.get("Project contingency"), 0.0) if factive.get("Project contingency", True) else 0.0
    proc_cont = fnum(fparam.get("Process contingency"), 0.0) if factive.get("Process contingency", True) else 0.0
    financial_land = fnum(fparam.get("Land"), 0.0) if factive.get("Land", True) else 0.0
    land += financial_land

    engineering = installed_total * eng_factor
    project_contingency = (installed_total + engineering) * proj_cont
    process_contingency = installed_total * proc_cont
    fci = installed_total + engineering + project_contingency + process_contingency
    simple_wc = fci * wc_factor_simple
    simple_startup = fci * startup_factor_simple
    tci = fci + simple_wc + simple_startup + financial_land

    # Materials
    material_rows = rd.rows_by_header("Materials")
    material_records = []
    material_opex = 0.0
    reference_material_costs = {}
    for row in material_rows:
        name = str(row.get("Material Name") or "Material")
        cost = annual_cost(
            fnum(row.get("Flow Rate")), str(row.get("Flow Unit") or ""),
            fnum(row.get("Linked Price")), str(row.get("Price Unit") or ""), hours
        )
        reference_material_costs[name.lower()] = cost
        if yes(row.get("Active?")):
            material_opex += cost
        material_records.append({
            "Material": name, "Active": yes(row.get("Active?")),
            "Annual Cost": cost, "Source ID": row.get("Source ID")
        })
    material_df = pd.DataFrame(material_records)

    # Utilities
    utility_rows = rd.rows_by_header("Utilities")
    utility_records = []
    electricity_opex = 0.0
    other_utility_opex = 0.0
    reference_utility_costs = {}
    for row in utility_rows:
        name = str(row.get("Utility Name") or "Utility")
        cost = annual_cost(
            fnum(row.get("Demand")), str(row.get("Unit") or ""),
            fnum(row.get("Linked Price")), str(row.get("Price Unit") or ""), hours
        )
        reference_utility_costs[name.lower()] = cost
        category = str(row.get("Category") or name).lower()
        if yes(row.get("Active?")):
            if "electric" in category or "electric" in name.lower():
                electricity_opex += cost
            else:
                other_utility_opex += cost
        utility_records.append({
            "Utility": name, "Active": yes(row.get("Active?")),
            "Annual Cost": cost, "Source ID": row.get("Source ID")
        })
    utility_df = pd.DataFrame(utility_records)

    # Waste
    waste_rows = rd.rows_by_header("Waste")
    waste_records = []
    waste_opex = 0.0
    for row in waste_rows:
        name = str(row.get("Waste Name") or "Waste")
        cost = annual_cost(
            fnum(row.get("Flow Rate")), str(row.get("Unit") or ""),
            fnum(row.get("Linked Disposal Cost")), str(row.get("Cost Unit") or ""), hours
        )
        if yes(row.get("Active?")):
            waste_opex += cost
        waste_records.append({
            "Waste": name, "Active": yes(row.get("Active?")),
            "Annual Cost": cost, "Source ID": row.get("Source ID")
        })
    waste_df = pd.DataFrame(waste_records)

    # Fixed OPEX sheet: can also carry published aggregate variable OPEX.
    fixed_rows = rd.rows_by_header("Fixed_OPEX")
    fixed_records = []
    fixed_opex = 0.0
    direct_variable_opex = 0.0

    labor_total = 0.0
    # First pass direct labor / direct annual
    for row in fixed_rows:
        if not yes(row.get("Active?")):
            continue
        basis = str(row.get("Basis") or "").strip().lower()
        category = str(row.get("Category") or "").strip().lower()
        value = fnum(row.get("Value / Factor"))
        item = str(row.get("Cost Item") or "")
        if basis == "direct annual" and "labor" in category:
            labor_total += value
        elif basis == "direct annual" and "labor" in item.lower():
            labor_total += value

    for row in fixed_rows:
        if not yes(row.get("Active?")):
            continue
        item = str(row.get("Cost Item") or "Fixed OPEX")
        basis = str(row.get("Basis") or "").strip().lower()
        category = str(row.get("Category") or "").strip().lower()
        value = fnum(row.get("Value / Factor"))
        if basis == "direct annual":
            annual = value
        elif basis == "fraction of installed":
            annual = value * installed_total
        elif basis == "fraction of fci":
            annual = value * fci
        elif basis == "fraction of labor":
            annual = value * labor_total
        else:
            annual = value

        if "variable" in category:
            direct_variable_opex += annual
        else:
            fixed_opex += annual
        fixed_records.append({
            "Cost Item": item, "Basis": row.get("Basis"), "Category": row.get("Category"),
            "Annual Cost": annual, "Source ID": row.get("Source ID")
        })
    fixed_rows_df = pd.DataFrame(fixed_records)

    # Detailed finance
    dparam = {}
    dactive = {}
    for row in detailed_rows:
        p = str(row.get("Parameter") or "").strip()
        if p:
            dparam[p] = row.get("Value")
            dactive[p] = yes(row.get("Active?"))

    finance_mode = str(dparam.get("Finance mode") or "Simple")
    detailed = {
        "construction_1": fnum(dparam.get("Construction year 1 share"), 1.0),
        "construction_2": fnum(dparam.get("Construction year 2 share"), 0.0),
        "construction_3": fnum(dparam.get("Construction year 3 share"), 0.0),
        "startup_revenue": fnum(dparam.get("Startup revenue factor"), 1.0),
        "startup_variable": fnum(dparam.get("Startup variable OPEX factor"), 1.0),
        "startup_fixed": fnum(dparam.get("Startup fixed OPEX factor"), 1.0),
        "working_capital_rate": fnum(dparam.get("Working capital rate"), 0.0),
        "replacement_npv": fnum(dparam.get("Replacement capital NPV"), 0.0),
        "replacement_interval": fnum(dparam.get("Replacement interval"), 0.0),
        "replacement_fraction": fnum(dparam.get("Replacement cost fraction"), 0.0),
        "decommissioning": fnum(dparam.get("Decommissioning factor"), 0.0),
        "salvage": fnum(dparam.get("Salvage factor"), 0.0),
        "equity_fraction": fnum(dparam.get("Equity fraction"), 1.0),
        "debt_fraction": fnum(dparam.get("Debt fraction"), 0.0),
        "debt_interest": fnum(dparam.get("Debt interest rate"), 0.0),
        "degradation_rate": fnum(dparam.get("Equipment degradation rate"), 0.0),
    }

    warnings = []
    if detailed["debt_fraction"] > 0:
        warnings.append(
            f"Debt financing ({detailed['debt_fraction']:.0%} debt at "
            f"{detailed['debt_interest']:.2%}) is present in the workbook but is not applied in V0.6."
        )
    if detailed["degradation_rate"] > 0:
        warnings.append(
            f"Equipment degradation ({detailed['degradation_rate']}) is present in the workbook "
            "but is not applied in V0.6."
        )

    ctx = ModelContext(
        case_name=case_name,
        process_category=process_category,
        primary_product=primary_product,
        hours=hours,
        life=life,
        currency=currency,
        reference_year=reference_year,
        product_rate=product_rate,
        product_unit=product_unit,
        selling_price=selling_price,
        selling_price_unit=selling_price_unit,
        equipment_df=equipment_df,
        material_df=material_df,
        utility_df=utility_df,
        waste_df=waste_df,
        fixed_rows_df=fixed_rows_df,
        purchased_equipment=purchased_total,
        installed_equipment=installed_total,
        fci=fci,
        tci=tci,
        land=land,
        material_opex=material_opex,
        electricity_opex=electricity_opex,
        other_utility_opex=other_utility_opex,
        waste_opex=waste_opex,
        direct_variable_opex=direct_variable_opex,
        fixed_opex=fixed_opex,
        target_return=target_return,
        tax_rate=tax_rate,
        depreciation_method=depreciation_method,
        depreciation_period=depreciation_period,
        wc_factor_simple=wc_factor_simple,
        startup_factor_simple=startup_factor_simple,
        salvage_factor_simple=salvage_factor_simple,
        finance_mode=finance_mode,
        detailed=detailed,
        warnings=warnings,
        reference_material_costs=reference_material_costs,
        reference_utility_costs=reference_utility_costs,
        reference_equipment_installed=reference_equipment_installed,
    )

    decision_rows = parse_decision_analysis(rd)
    validation_df = pd.DataFrame(validation_rows)
    return ctx, rd, decision_rows, validation_df


def parse_decision_analysis(rd: WorkbookReader) -> pd.DataFrame:
    if not rd.has_sheet("Decision_Analysis"):
        return pd.DataFrame()
    ws = rd.wb["Decision_Analysis"]
    records = []
    # Scan headers dynamically.
    for r in range(1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if "Parameter" in vals and ("Low Factor" in vals or "Lower Bound" in vals):
            headers = [str(v).strip() if v is not None else "" for v in vals]
            section = "Sensitivity" if "Low Factor" in headers else "BreakEven"
            rr = r + 1
            while rr <= ws.max_row:
                rowvals = [ws.cell(rr, c).value for c in range(1, ws.max_column + 1)]
                if all(v in (None, "") for v in rowvals):
                    break
                if str(rowvals[0] or "").strip().startswith(("B.", "C.")):
                    break
                rec = {"Section": section}
                for h, v in zip(headers, rowvals):
                    if h:
                        rec[h] = v
                records.append(rec)
                rr += 1
        if "Scenario" in vals and "Equipment CAPEX" in vals:
            headers = [str(v).strip() if v is not None else "" for v in vals]
            rr = r + 1
            while rr <= ws.max_row:
                rowvals = [ws.cell(rr, c).value for c in range(1, ws.max_column + 1)]
                if all(v in (None, "") for v in rowvals):
                    break
                if str(rowvals[0] or "").strip().startswith("C."):
                    break
                rec = {"Section": "Scenario"}
                for h, v in zip(headers, rowvals):
                    if h:
                        rec[h] = v
                records.append(rec)
                rr += 1
    return pd.DataFrame(records)


# ---------------------------------------------------------------------
# Core operating quantities
# ---------------------------------------------------------------------

def annual_product_quantity(ctx: ModelContext, hours_factor=1.0, output_factor=1.0) -> float:
    unit = ctx.product_unit.strip().lower().replace("³", "3")
    rate = ctx.product_rate
    if unit in {"kg/h", "nm3/h", "m3/h", "kwh/h"}:
        return rate * ctx.hours * hours_factor * output_factor
    if unit in {"kg/year", "ton/year", "m3/year", "per year"}:
        return rate * output_factor
    if unit == "ton/h":
        return rate * ctx.hours * hours_factor * output_factor
    return rate * ctx.hours * hours_factor * output_factor


def opex_components(ctx: ModelContext, hours_factor=1.0, product_output_factor=1.0,
                    material_factor=1.0, utility_factor=1.0,
                    electricity_factor=1.0, electricity_demand_factor=1.0,
                    fixed_factor=1.0, direct_variable_factor=1.0) -> Dict[str, float]:
    # Active flow-based costs scale with operating hours.
    return {
        "Materials": ctx.material_opex * hours_factor * material_factor,
        "Electricity": ctx.electricity_opex * hours_factor * utility_factor * electricity_factor * electricity_demand_factor,
        "Other utilities": ctx.other_utility_opex * hours_factor * utility_factor,
        "Waste": ctx.waste_opex * hours_factor,
        "Direct variable OPEX": ctx.direct_variable_opex * hours_factor * direct_variable_factor,
        "Fixed OPEX": ctx.fixed_opex * fixed_factor,
    }


def total_opex(components: Dict[str, float]) -> float:
    return float(sum(components.values()))


# ---------------------------------------------------------------------
# Cash flow / finance
# ---------------------------------------------------------------------

def depreciation_schedule(ctx: ModelContext, depreciable_capital: float) -> List[float]:
    method = ctx.depreciation_method.lower()
    if "macrs" in method:
        sched = [depreciable_capital * r for r in MACRS20]
        return sched + [0.0] * max(0, ctx.life - len(sched))
    period = max(1, ctx.depreciation_period)
    sched = [depreciable_capital / period] * min(period, ctx.life)
    return sched + [0.0] * max(0, ctx.life - len(sched))


def build_cashflows(
    ctx: ModelContext,
    price: float,
    *,
    tci_factor=1.0,
    hours_factor=1.0,
    product_output_factor=1.0,
    material_factor=1.0,
    utility_factor=1.0,
    electricity_factor=1.0,
    electricity_demand_factor=1.0,
    fixed_factor=1.0,
    direct_variable_factor=1.0,
    tci_adjustment=0.0,
    variable_opex_adjustment=0.0,
    replacement_fraction_factor=1.0,
    replacement_interval_factor=1.0,
) -> pd.DataFrame:

    tci = ctx.tci * tci_factor + tci_adjustment
    land_scaled = ctx.land * tci_factor
    depreciable = max(0.0, tci - land_scaled)

    annual_product = annual_product_quantity(ctx, hours_factor, product_output_factor)
    comps = opex_components(
        ctx,
        hours_factor=hours_factor,
        product_output_factor=product_output_factor,
        material_factor=material_factor,
        utility_factor=utility_factor,
        electricity_factor=electricity_factor,
        electricity_demand_factor=electricity_demand_factor,
        fixed_factor=fixed_factor,
        direct_variable_factor=direct_variable_factor,
    )
    # Case-specific sensitivity adjustments can add/subtract annual variable OPEX.
    comps["Direct variable OPEX"] += variable_opex_adjustment
    variable_opex = (
        comps["Materials"] + comps["Electricity"] + comps["Other utilities"]
        + comps["Waste"] + comps["Direct variable OPEX"]
    )
    fixed_opex = comps["Fixed OPEX"]

    detailed_mode = "detailed" in ctx.finance_mode.lower() or "h2a" in ctx.finance_mode.lower()

    rows = []
    if not detailed_mode:
        dep = depreciation_schedule(ctx, max(0.0, ctx.fci * tci_factor))
        initial = tci
        rows.append({
            "Time": 0, "Phase": "Initial capital", "Operating Year": 0,
            "Revenue": 0.0, "Fixed OPEX": 0.0, "Variable OPEX": 0.0,
            "Depreciation": 0.0, "Tax": 0.0, "Working Capital Change": 0.0,
            "Replacement Capital": 0.0, "Terminal Adjustment": 0.0,
            "Net Cash Flow": -initial,
        })
        for year in range(1, ctx.life + 1):
            revenue = annual_product * price
            depreciation = dep[year - 1] if year - 1 < len(dep) else 0.0
            taxable = max(0.0, revenue - fixed_opex - variable_opex - depreciation)
            tax = taxable * ctx.tax_rate
            terminal = 0.0
            if year == ctx.life:
                terminal += ctx.fci * tci_factor * ctx.salvage_factor_simple
                terminal += ctx.fci * tci_factor * ctx.wc_factor_simple
            net = revenue - fixed_opex - variable_opex - tax + terminal
            rows.append({
                "Time": year, "Phase": "Operation", "Operating Year": year,
                "Revenue": revenue, "Fixed OPEX": fixed_opex, "Variable OPEX": variable_opex,
                "Depreciation": depreciation, "Tax": tax, "Working Capital Change": 0.0,
                "Replacement Capital": 0.0, "Terminal Adjustment": terminal,
                "Net Cash Flow": net,
            })
    else:
        d = ctx.detailed
        fractions = [d["construction_1"], d["construction_2"], d["construction_3"]]
        fractions = [x for x in fractions if x > 0]
        if not fractions:
            fractions = [1.0]
        s = sum(fractions)
        fractions = [x / s for x in fractions]

        replacement_pv = d["replacement_npv"] * tci_factor
        for i, frac in enumerate(fractions):
            replacement_pv_here = replacement_pv if i == 0 and replacement_pv > 0 else 0.0
            rows.append({
                "Time": i, "Phase": "Construction", "Operating Year": 0,
                "Revenue": 0.0, "Fixed OPEX": 0.0, "Variable OPEX": 0.0,
                "Depreciation": 0.0, "Tax": 0.0, "Working Capital Change": 0.0,
                "Replacement Capital": replacement_pv_here, "Terminal Adjustment": 0.0,
                "Net Cash Flow": -tci * frac - replacement_pv_here,
            })

        dep = depreciation_schedule(ctx, depreciable)
        wc_balance = 0.0

        interval = d["replacement_interval"] * replacement_interval_factor
        repl_fraction = d["replacement_fraction"] * replacement_fraction_factor
        replacement_years = set()
        if replacement_pv <= 0 and interval > 0 and repl_fraction > 0:
            y = int(round(interval))
            if y > 0:
                replacement_years = set(range(y, ctx.life, y))

        for year in range(1, ctx.life + 1):
            startup = year == 1
            rev_factor = d["startup_revenue"] if startup else 1.0
            fixed_start = d["startup_fixed"] if startup else 1.0
            var_start = d["startup_variable"] if startup else 1.0

            revenue = annual_product * price * rev_factor
            fixed_y = fixed_opex * fixed_start
            variable_y = variable_opex * var_start
            depreciation = dep[year - 1] if year - 1 < len(dep) else 0.0

            taxable = max(0.0, revenue - fixed_y - variable_y - depreciation)
            tax = taxable * ctx.tax_rate

            target_wc = d["working_capital_rate"] * (fixed_y + variable_y)
            wc_change = target_wc - wc_balance
            wc_balance = target_wc

            replacement = repl_fraction * tci if year in replacement_years else 0.0

            terminal = 0.0
            if year == ctx.life:
                terminal += d["salvage"] * tci
                terminal -= d["decommissioning"] * depreciable
                terminal += wc_balance

            net = revenue - fixed_y - variable_y - tax - wc_change - replacement + terminal
            rows.append({
                "Time": len(fractions) + year - 1,
                "Phase": "Startup" if startup else "Operation",
                "Operating Year": year,
                "Revenue": revenue, "Fixed OPEX": fixed_y, "Variable OPEX": variable_y,
                "Depreciation": depreciation, "Tax": tax, "Working Capital Change": wc_change,
                "Replacement Capital": replacement, "Terminal Adjustment": terminal,
                "Net Cash Flow": net,
            })

    df = pd.DataFrame(rows)
    df["Cumulative Cash Flow"] = df["Net Cash Flow"].cumsum()
    return df



def fast_cashflow_values(
    ctx: ModelContext,
    price: float,
    *,
    tci_factor=1.0,
    hours_factor=1.0,
    product_output_factor=1.0,
    material_factor=1.0,
    utility_factor=1.0,
    electricity_factor=1.0,
    electricity_demand_factor=1.0,
    fixed_factor=1.0,
    direct_variable_factor=1.0,
    tci_adjustment=0.0,
    variable_opex_adjustment=0.0,
    replacement_fraction_factor=1.0,
    replacement_interval_factor=1.0,
) -> List[float]:
    """Fast numeric-only equivalent of build_cashflows for repeated root solves."""
    tci = ctx.tci * tci_factor + tci_adjustment
    land_scaled = ctx.land * tci_factor
    depreciable = max(0.0, tci - land_scaled)

    annual_product = annual_product_quantity(ctx, hours_factor, product_output_factor)
    comps = opex_components(
        ctx,
        hours_factor=hours_factor,
        product_output_factor=product_output_factor,
        material_factor=material_factor,
        utility_factor=utility_factor,
        electricity_factor=electricity_factor,
        electricity_demand_factor=electricity_demand_factor,
        fixed_factor=fixed_factor,
        direct_variable_factor=direct_variable_factor,
    )
    comps["Direct variable OPEX"] += variable_opex_adjustment
    variable_opex = (
        comps["Materials"] + comps["Electricity"] + comps["Other utilities"]
        + comps["Waste"] + comps["Direct variable OPEX"]
    )
    fixed_opex = comps["Fixed OPEX"]

    detailed_mode = "detailed" in ctx.finance_mode.lower() or "h2a" in ctx.finance_mode.lower()

    if not detailed_mode:
        dep = depreciation_schedule(ctx, max(0.0, ctx.fci * tci_factor))
        cfs = [-tci]
        for year in range(1, ctx.life + 1):
            revenue = annual_product * price
            depreciation = dep[year - 1] if year - 1 < len(dep) else 0.0
            taxable = max(0.0, revenue - fixed_opex - variable_opex - depreciation)
            tax = taxable * ctx.tax_rate
            terminal = 0.0
            if year == ctx.life:
                terminal += ctx.fci * tci_factor * ctx.salvage_factor_simple
                terminal += ctx.fci * tci_factor * ctx.wc_factor_simple
            cfs.append(revenue - fixed_opex - variable_opex - tax + terminal)
        return cfs

    d = ctx.detailed
    fractions = [d["construction_1"], d["construction_2"], d["construction_3"]]
    fractions = [x for x in fractions if x > 0] or [1.0]
    total_frac = sum(fractions)
    fractions = [x / total_frac for x in fractions]

    replacement_pv = d["replacement_npv"] * tci_factor
    cfs = []
    for i, frac in enumerate(fractions):
        replacement_here = replacement_pv if i == 0 and replacement_pv > 0 else 0.0
        cfs.append(-tci * frac - replacement_here)

    dep = depreciation_schedule(ctx, depreciable)
    wc_balance = 0.0

    interval = d["replacement_interval"] * replacement_interval_factor
    repl_fraction = d["replacement_fraction"] * replacement_fraction_factor
    replacement_years = set()
    if replacement_pv <= 0 and interval > 0 and repl_fraction > 0:
        step = int(round(interval))
        if step > 0:
            replacement_years = set(range(step, ctx.life, step))

    for year in range(1, ctx.life + 1):
        startup = year == 1
        revenue = annual_product * price * (d["startup_revenue"] if startup else 1.0)
        fixed_y = fixed_opex * (d["startup_fixed"] if startup else 1.0)
        variable_y = variable_opex * (d["startup_variable"] if startup else 1.0)
        depreciation = dep[year - 1] if year - 1 < len(dep) else 0.0

        taxable = max(0.0, revenue - fixed_y - variable_y - depreciation)
        tax = taxable * ctx.tax_rate

        target_wc = d["working_capital_rate"] * (fixed_y + variable_y)
        wc_change = target_wc - wc_balance
        wc_balance = target_wc

        replacement = repl_fraction * tci if year in replacement_years else 0.0
        terminal = 0.0
        if year == ctx.life:
            terminal += d["salvage"] * tci
            terminal -= d["decommissioning"] * depreciable
            terminal += wc_balance

        cfs.append(revenue - fixed_y - variable_y - tax - wc_change - replacement + terminal)

    return cfs


def evaluate_case(ctx: ModelContext, **mods) -> Dict[str, object]:
    price = mods.pop("price", ctx.selling_price)

    def builder(p):
        return fast_cashflow_values(ctx, p, **mods)

    cashflows = fast_cashflow_values(ctx, price, **mods)
    cf_at_price = build_cashflows(ctx, price, **mods)
    model_npv = npv(cashflows, ctx.target_return)
    irr = solve_irr(cashflows)
    msp = solve_root_price(builder, ctx.target_return)

    payback = np.nan
    cumulative = cf_at_price["Cumulative Cash Flow"].tolist()
    times = cf_at_price["Time"].tolist()
    for i, value in enumerate(cumulative):
        if value >= 0:
            if i == 0:
                payback = float(times[i])
            else:
                prev = cumulative[i - 1]
                cur = value
                span = float(times[i] - times[i - 1])
                frac = (-prev) / (cur - prev) if cur != prev else 0.0
                payback = float(times[i - 1]) + frac * span
            break

    # Annual nominal quantities outside startup effects for screening outputs.
    annual_product = annual_product_quantity(
        ctx,
        mods.get("hours_factor", 1.0),
        mods.get("product_output_factor", 1.0),
    )
    comps = opex_components(
        ctx,
        hours_factor=mods.get("hours_factor", 1.0),
        product_output_factor=mods.get("product_output_factor", 1.0),
        material_factor=mods.get("material_factor", 1.0),
        utility_factor=mods.get("utility_factor", 1.0),
        electricity_factor=mods.get("electricity_factor", 1.0),
        electricity_demand_factor=mods.get("electricity_demand_factor", 1.0),
        fixed_factor=mods.get("fixed_factor", 1.0),
        direct_variable_factor=mods.get("direct_variable_factor", 1.0),
    )
    comps["Direct variable OPEX"] += mods.get("variable_opex_adjustment", 0.0)
    annual_opex = total_opex(comps)
    annual_revenue = annual_product * price

    # Simple BCR at the project discount rate.
    revenue_series = cf_at_price["Revenue"].tolist()
    cost_series = (
        cf_at_price["Fixed OPEX"] + cf_at_price["Variable OPEX"] + cf_at_price["Tax"]
        + cf_at_price["Working Capital Change"] + cf_at_price["Replacement Capital"]
    ).tolist()
    # Capital outflows embedded in negative NCF construction rows are added to costs.
    for i, phase in enumerate(cf_at_price["Phase"].tolist()):
        if phase in {"Construction", "Initial capital"}:
            cost_series[i] += abs(cf_at_price["Net Cash Flow"].iloc[i])
    pv_rev = sum(v / ((1 + ctx.target_return) ** t) for t, v in enumerate(revenue_series))
    pv_cost = sum(max(0.0, v) / ((1 + ctx.target_return) ** t) for t, v in enumerate(cost_series))
    bcr = safe_div(pv_rev, pv_cost)

    tci_calc = ctx.tci * mods.get("tci_factor", 1.0) + mods.get("tci_adjustment", 0.0)

    return {
        "Price": price,
        "Annual Product": annual_product,
        "Annual OPEX": annual_opex,
        "Annual Revenue": annual_revenue,
        "TCI": tci_calc,
        "NPV": model_npv,
        "IRR": irr,
        "Payback": payback,
        "MSP": msp,
        "BCR": bcr,
        "Cash Flow": cf_at_price,
        "OPEX Components": comps,
    }


# ---------------------------------------------------------------------
# Sensitivity / scenarios / break-even / scale
# ---------------------------------------------------------------------

def find_reference_cost(d: Dict[str, float], keyword: str) -> float:
    keyword = keyword.lower()
    return sum(v for k, v in d.items() if keyword in k.lower())

def parameter_mods(ctx: ModelContext, parameter: str, factor: float) -> Dict[str, float]:
    p = parameter.strip().lower()
    mods = {}

    if p in {"equipment capex", "total capital investment"}:
        mods["tci_factor"] = factor
    elif p == "material prices":
        mods["material_factor"] = factor
    elif p == "utility prices":
        mods["utility_factor"] = factor
    elif p == "product selling price":
        mods["price"] = ctx.selling_price * factor
    elif p in {"operating hours", "operating capacity factor", "capacity factor"}:
        mods["hours_factor"] = factor
    elif p == "primary product output":
        mods["product_output_factor"] = factor
    elif p == "electricity demand":
        mods["electricity_demand_factor"] = factor
    elif p in {"fixed opex", "total fixed operating cost"}:
        mods["fixed_factor"] = factor

    # Published-case decomposition logic.
    elif p == "biomass feedstock price":
        ref = find_reference_cost(ctx.reference_material_costs, "biomass")
        mods["variable_opex_adjustment"] = ref * (factor - 1.0)
    elif p == "plant efficiency":
        # Higher efficiency means lower feedstock requirement for same product.
        ref = find_reference_cost(ctx.reference_material_costs, "biomass")
        if factor > 0:
            mods["variable_opex_adjustment"] = ref * ((1.0 / factor) - 1.0)
    elif p == "labor requirement":
        # V0.6 lacks a fully disaggregated labor model for every case.
        # Use the published fixed-cost block as the screening proxy.
        mods["fixed_factor"] = factor
    elif p == "electricity price":
        ref = find_reference_cost(ctx.reference_utility_costs, "electric")
        if ref > 0 and ctx.direct_variable_opex > 0:
            mods["variable_opex_adjustment"] = ref * (factor - 1.0)
        else:
            mods["electricity_factor"] = factor
    elif p == "stack electrical usage":
        ref = find_reference_cost(ctx.reference_utility_costs, "electric")
        if ref > 0 and ctx.direct_variable_opex > 0:
            mods["variable_opex_adjustment"] = ref * (factor - 1.0)
        else:
            mods["electricity_demand_factor"] = factor
    elif p == "stack cost":
        stack_installed = find_reference_cost(ctx.reference_equipment_installed, "stack")
        mods["tci_adjustment"] = stack_installed * (factor - 1.0)
    elif p == "stack replacement cost":
        mods["replacement_fraction_factor"] = factor
    elif p == "stack replacement interval":
        mods["replacement_interval_factor"] = factor

    return mods


def run_sensitivity(ctx: ModelContext, decision_df: pd.DataFrame) -> pd.DataFrame:
    if decision_df.empty:
        return pd.DataFrame()
    rows = decision_df[decision_df["Section"].eq("Sensitivity")] if "Section" in decision_df.columns else pd.DataFrame()
    out = []
    baseline = evaluate_case(ctx)
    for _, row in rows.iterrows():
        if not yes(row.get("Analyze?")):
            continue
        parameter = str(row.get("Parameter") or "")
        low_factor = fnum(row.get("Low Factor"), 1.0)
        high_factor = fnum(row.get("High Factor"), 1.0)

        low = evaluate_case(ctx, **parameter_mods(ctx, parameter, low_factor))
        high = evaluate_case(ctx, **parameter_mods(ctx, parameter, high_factor))
        impact = max(abs(low["MSP"] - baseline["MSP"]), abs(high["MSP"] - baseline["MSP"]))
        out.append({
            "Parameter": parameter,
            "Low Factor": low_factor,
            "High Factor": high_factor,
            "Low NPV": low["NPV"],
            "High NPV": high["NPV"],
            "Low MSP": low["MSP"],
            "Baseline MSP": baseline["MSP"],
            "High MSP": high["MSP"],
            "Max |MSP Change|": impact,
        })
    if not out:
        return pd.DataFrame()
    df = pd.DataFrame(out).sort_values("Max |MSP Change|", ascending=False).reset_index(drop=True)
    max_imp = df["Max |MSP Change|"].max()
    df["Impact Score"] = df["Max |MSP Change|"] / max_imp if max_imp > 0 else 0.0
    df["Priority"] = df["Impact Score"].apply(lambda x: "High" if x >= 0.67 else ("Medium" if x >= 0.33 else "Low"))
    df["Rank"] = np.arange(1, len(df) + 1)
    return df


def run_scenarios(ctx: ModelContext, decision_df: pd.DataFrame) -> pd.DataFrame:
    if decision_df.empty or "Section" not in decision_df.columns:
        return pd.DataFrame()
    rows = decision_df[decision_df["Section"].eq("Scenario")]
    out = []
    for _, row in rows.iterrows():
        name = str(row.get("Scenario") or "Scenario")
        mods = {
            "tci_factor": fnum(row.get("Equipment CAPEX"), 1.0),
            "material_factor": fnum(row.get("Material Price"), 1.0),
            "utility_factor": fnum(row.get("Utility Price"), 1.0),
            "fixed_factor": fnum(row.get("Fixed OPEX"), 1.0),
            "price": ctx.selling_price * fnum(row.get("Product Price"), 1.0),
            "hours_factor": fnum(row.get("Operating Hours / CF"), 1.0),
            "product_output_factor": fnum(row.get("Product Output"), 1.0),
        }
        res = evaluate_case(ctx, **mods)
        status = "Viable" if res["NPV"] >= 0 and (np.isnan(res["IRR"]) or res["IRR"] >= ctx.target_return) else "Economically weak"
        out.append({
            "Scenario": name,
            "NPV": res["NPV"], "IRR": res["IRR"], "MSP": res["MSP"],
            "Annual OPEX": res["Annual OPEX"], "Annual Revenue": res["Annual Revenue"],
            "Status": status,
        })
    return pd.DataFrame(out)


def solve_break_even_factor(ctx: ModelContext, parameter: str, lo: float, hi: float) -> Tuple[float, str]:
    def val(f):
        return evaluate_case(ctx, **parameter_mods(ctx, parameter, f))["NPV"]
    vlo, vhi = val(lo), val(hi)
    if np.isnan(vlo) or np.isnan(vhi) or vlo * vhi > 0:
        return np.nan, "Not bracketed"
    for _ in range(80):
        mid = (lo + hi) / 2.0
        vmid = val(mid)
        if abs(vmid) < 1e-5:
            return mid, "Solved"
        if vlo * vmid <= 0:
            hi = mid
            vhi = vmid
        else:
            lo = mid
            vlo = vmid
    return (lo + hi) / 2.0, "Solved"


def run_break_even(ctx: ModelContext, decision_df: pd.DataFrame) -> pd.DataFrame:
    if decision_df.empty or "Section" not in decision_df.columns:
        return pd.DataFrame()
    rows = decision_df[decision_df["Section"].eq("BreakEven")]
    out = []
    for _, row in rows.iterrows():
        if not yes(row.get("Analyze?")):
            continue
        parameter = str(row.get("Parameter") or "")
        lo = fnum(row.get("Lower Bound"), 0.1)
        hi = fnum(row.get("Upper Bound"), 5.0)
        factor, status = solve_break_even_factor(ctx, parameter, lo, hi)
        out.append({
            "Parameter": parameter, "Lower Bound": lo, "Upper Bound": hi,
            "Break-even Factor": factor, "Change vs Base (%)": (factor - 1.0) * 100 if not np.isnan(factor) else np.nan,
            "Status": status,
        })
    return pd.DataFrame(out)


def run_scale(ctx: ModelContext, decision_df: pd.DataFrame) -> pd.DataFrame:
    if decision_df.empty or "Section" not in decision_df.columns:
        return pd.DataFrame()
    rows = decision_df[decision_df["Section"].eq("BreakEven")]
    factors = []
    for _, row in rows.iterrows():
        if yes(row.get("Scale Active?")) and row.get("Scale Factor") not in (None, ""):
            f = fnum(row.get("Scale Factor"), np.nan)
            if not np.isnan(f):
                factors.append(f)
    factors = sorted(set(factors))
    out = []
    for factor in factors:
        # Rebuild equipment CAPEX using each equipment's own exponent.
        if ctx.equipment_df.empty:
            cap_factor = factor
        else:
            base_inst = ctx.installed_equipment
            scaled_inst = float(
                (ctx.equipment_df["Base Cost"]
                 * ((ctx.equipment_df["Actual Capacity"] * factor) / ctx.equipment_df["Base Capacity"]) ** ctx.equipment_df["Scaling Exponent"]
                 * ctx.equipment_df["Installation Factor"]).sum()
            )
            cap_factor = safe_div(scaled_inst, base_inst, factor)

        res = evaluate_case(
            ctx,
            tci_factor=cap_factor,
            hours_factor=1.0,
            product_output_factor=factor,
            material_factor=factor,
            utility_factor=factor,
            direct_variable_factor=factor,
        )
        out.append({
            "Scale Factor": factor, "Effective Capital Factor": cap_factor,
            "TCI": res["TCI"], "Annual Product": res["Annual Product"],
            "NPV": res["NPV"], "IRR": res["IRR"], "MSP": res["MSP"],
            "Status": "Viable" if res["NPV"] >= 0 else "Economically weak",
        })
    return pd.DataFrame(out)


# ---------------------------------------------------------------------
# Environmental screening
# ---------------------------------------------------------------------

def calculate_environmental(ctx: ModelContext, rd: WorkbookReader) -> pd.DataFrame:
    rows = rd.rows_by_header("Environmental")
    active = {str(r.get("Parameter") or "").strip().lower(): r for r in rows if yes(r.get("Active?"))}
    if not active:
        return pd.DataFrame()

    def env_value(keyword, default=0.0):
        for k, r in active.items():
            if keyword in k:
                return fnum(r.get("Value"), default)
        return default

    ef_elec = env_value("electricity ef")
    ef_heat = env_value("heat ef")
    ef_feed = env_value("feedstock/material ef")
    direct_h = env_value("direct process emissions")
    ref_ci = env_value("reference carbon intensity")
    resource_eff = env_value("resource efficiency")
    ref_re = env_value("reference resource efficiency")

    # Active annual flow quantities
    electricity_kwh = 0.0
    heat_gj = 0.0
    for _, r in ctx.utility_df.iterrows():
        if not r["Active"]:
            continue
        name = str(r["Utility"]).lower()
        # Use the reference workbook rows for quantities via source costs is not robust,
        # so re-read utility sheet from rd.
    util_rows = rd.rows_by_header("Utilities")
    for row in util_rows:
        if not yes(row.get("Active?")):
            continue
        qty, unit = annual_quantity(fnum(row.get("Demand")), str(row.get("Unit") or ""), ctx.hours)
        name = str(row.get("Utility Name") or "").lower()
        if "electric" in name:
            if unit.lower() == "kwh":
                electricity_kwh += qty
            elif unit.lower() == "mwh":
                electricity_kwh += qty * 1000
        if "heat" in name:
            if unit.lower() == "gj":
                heat_gj += qty
            elif unit.lower() == "mj":
                heat_gj += qty / 1000

    feed_kg = 0.0
    for row in rd.rows_by_header("Materials"):
        if not yes(row.get("Active?")):
            continue
        if str(row.get("Role") or "").lower() != "feedstock":
            continue
        qty, unit = annual_quantity(fnum(row.get("Flow Rate")), str(row.get("Flow Unit") or ""), ctx.hours)
        feed_kg += qty * 1000 if unit.lower() == "ton" else qty

    e_elec = electricity_kwh * ef_elec
    e_heat = heat_gj * ef_heat
    e_feed = feed_kg * ef_feed
    e_direct = direct_h * ctx.hours
    total = e_elec + e_heat + e_feed + e_direct
    product = annual_product_quantity(ctx)
    ci = safe_div(total, product)
    saving = (ref_ci - ci) / ref_ci * 100 if ref_ci > 0 and not np.isnan(ci) else np.nan
    annual_saving = (ref_ci - ci) * product if ref_ci > 0 and not np.isnan(ci) else np.nan

    return pd.DataFrame([
        ["Annual electricity use", electricity_kwh, "kWh/year"],
        ["Annual heat use", heat_gj, "GJ/year"],
        ["Annual feedstock", feed_kg, "kg/year"],
        ["Electricity emissions", e_elec, "kg CO2e/year"],
        ["Heat emissions", e_heat, "kg CO2e/year"],
        ["Feedstock emissions", e_feed, "kg CO2e/year"],
        ["Direct emissions", e_direct, "kg CO2e/year"],
        ["Total screening emissions", total, "kg CO2e/year"],
        ["Carbon intensity", ci, "kg CO2e/product unit"],
        ["Carbon saving vs reference", saving, "%"],
        ["Annual carbon saving", annual_saving, "kg CO2e/year"],
        ["Resource efficiency", resource_eff, "fraction"],
        ["Reference resource efficiency", ref_re, "fraction"],
        ["Resource-efficiency improvement", (resource_eff-ref_re)*100 if resource_eff or ref_re else np.nan, "percentage points"],
    ], columns=["Indicator", "Value", "Unit"])


# ---------------------------------------------------------------------
# Validation comparison
# ---------------------------------------------------------------------

def validation_comparison(
    ctx: ModelContext,
    validation_df: pd.DataFrame,
    baseline: Dict[str, object],
    sensitivity: pd.DataFrame,
) -> pd.DataFrame:
    if validation_df.empty or "Active?" not in validation_df.columns:
        return pd.DataFrame()
    rows = []
    fixed_per_product = safe_div(ctx.fixed_opex, baseline["Annual Product"])
    variable_per_product = safe_div(
        ctx.material_opex + ctx.electricity_opex + ctx.other_utility_opex
        + ctx.waste_opex + ctx.direct_variable_opex,
        baseline["Annual Product"]
    )
    electricity_ref = find_reference_cost(ctx.reference_utility_costs, "electric")
    electricity_per_product = safe_div(electricity_ref, baseline["Annual Product"])
    capital_contribution_proxy = baseline["MSP"] - fixed_per_product - electricity_per_product

    for _, row in validation_df.iterrows():
        if not yes(row.get("Active?")):
            continue
        metric = str(row.get("Metric") or "")
        published = fnum(row.get("Published Value"), np.nan)
        m = metric.lower()

        calculated = np.nan
        note = ""

        if "annual h2 production" in m:
            calculated = baseline["Annual Product"]
        elif "h2 design production" in m:
            # Use design on-stream daily output derived from product rate.
            if ctx.product_unit.lower() == "kg/h":
                calculated = ctx.product_rate * 24
        elif "total initial investment" in m:
            calculated = baseline["TCI"]
        elif "annual fixed opex" in m:
            calculated = ctx.fixed_opex
        elif "annual variable opex" in m:
            calculated = (
                ctx.material_opex + ctx.electricity_opex + ctx.other_utility_opex
                + ctx.waste_opex + ctx.direct_variable_opex
            )
        elif "hydrogen production cost" in m or "total h2 production cost" in m or "/ msp" in m:
            calculated = baseline["MSP"]
        elif "fixed o&m contribution" in m:
            calculated = fixed_per_product
        elif "electricity contribution" in m:
            calculated = electricity_per_product
        elif "capital cost contribution" in m:
            calculated = capital_contribution_proxy
            note = "Residual MSP contribution proxy, not a full H2A component decomposition."
        elif "electricity price sensitivity" in m and not sensitivity.empty:
            sr = sensitivity[sensitivity["Parameter"].str.lower().eq("electricity price")]
            if not sr.empty:
                calculated = sr.iloc[0]["Low MSP"] if "low" in m else sr.iloc[0]["High MSP"]
        elif "process energy efficiency" in m:
            # This is only available when the workbook explicitly provides enough
            # energy-basis data. V0.6 does not invent missing LHVs.
            note = "Not calculated: energy-basis/LHV inputs are not explicit in the universal workbook."

        deviation = pct_dev(calculated, published) if not np.isnan(calculated) else np.nan
        absdev = abs(deviation) if not np.isnan(deviation) else np.nan
        if np.isnan(absdev):
            status = "Not calculated"
        elif absdev <= 1:
            status = "PASS"
        elif absdev <= 5:
            status = "Screening match"
        else:
            status = "Gap"

        rows.append({
            "Metric": metric,
            "Published": published,
            "Calculated": calculated,
            "Deviation (%)": deviation,
            "Status": status,
            "Source ID": row.get("Source ID"),
            "Source Location": row.get("Location in Source"),
            "Acceptance Rule": row.get("Acceptance Rule"),
            "Notes": note or row.get("Notes"),
        })
    return pd.DataFrame(rows)


# ---------------------------------------------------------------------
# Charts
# ---------------------------------------------------------------------

def save_charts(
    ctx: ModelContext,
    out_dir: Path,
    stem: str,
    baseline: Dict[str, object],
    sensitivity: pd.DataFrame,
    scenarios: pd.DataFrame,
    validation: pd.DataFrame,
) -> List[str]:
    paths = []

    # CAPEX donut
    capex_path = out_dir / f"{stem}_CAPEX.png"
    if not ctx.equipment_df.empty and ctx.equipment_df["Installed Cost"].sum() > 0:
        data = ctx.equipment_df.groupby("Equipment")["Installed Cost"].sum().sort_values(ascending=False)
        fig, ax = plt.subplots(figsize=(8.5, 6.5))
        ax.pie(data.values, labels=data.index, autopct="%1.1f%%", startangle=90,
               wedgeprops={"width":0.42})
        ax.set_title("CAPEX breakdown — installed equipment / cost blocks")
        fig.tight_layout()
        fig.savefig(capex_path, dpi=220)
        plt.close(fig)
        paths.append(str(capex_path))

    # OPEX donut
    opex_path = out_dir / f"{stem}_OPEX.png"
    comps = pd.Series(baseline["OPEX Components"])
    comps = comps[comps > 0]
    if len(comps):
        fig, ax = plt.subplots(figsize=(8.5, 6.5))
        ax.pie(comps.values, labels=comps.index, autopct="%1.1f%%", startangle=90,
               wedgeprops={"width":0.42})
        ax.set_title("Annual OPEX breakdown")
        fig.tight_layout()
        fig.savefig(opex_path, dpi=220)
        plt.close(fig)
        paths.append(str(opex_path))

    # Cash flow
    cf_path = out_dir / f"{stem}_Cash_Flow.png"
    cf = baseline["Cash Flow"]
    fig, ax = plt.subplots(figsize=(9.5, 5.8))
    ax.plot(cf["Time"], cf["Cumulative Cash Flow"], marker="o", markersize=3)
    ax.axhline(0, linewidth=1)
    ax.set_xlabel("Project time")
    ax.set_ylabel(f"Cumulative cash flow ({ctx.currency})")
    ax.set_title("Cumulative project cash flow at entered selling price")
    fig.tight_layout()
    fig.savefig(cf_path, dpi=220)
    plt.close(fig)
    paths.append(str(cf_path))

    # Sensitivity
    if not sensitivity.empty:
        sens_path = out_dir / f"{stem}_Sensitivity.png"
        plot = sensitivity.sort_values("Max |MSP Change|", ascending=True)
        fig, ax = plt.subplots(figsize=(9.5, max(5.5, len(plot) * 0.55)))
        ax.barh(plot["Parameter"], plot["Max |MSP Change|"])
        ax.set_xlabel(f"Maximum absolute MSP change ({ctx.currency}/product unit)")
        ax.set_title("Sensitivity screening — MSP impact")
        fig.tight_layout()
        fig.savefig(sens_path, dpi=220)
        plt.close(fig)
        paths.append(str(sens_path))

    # Scenarios
    if not scenarios.empty:
        scen_path = out_dir / f"{stem}_Scenarios.png"
        fig, ax = plt.subplots(figsize=(8.8, 5.7))
        bars = ax.bar(scenarios["Scenario"], scenarios["NPV"])
        ax.axhline(0, linewidth=1)
        ax.set_ylabel(f"NPV ({ctx.currency})")
        ax.set_title("Scenario comparison")
        ax.tick_params(axis="x", rotation=20)
        fig.tight_layout()
        fig.savefig(scen_path, dpi=220)
        plt.close(fig)
        paths.append(str(scen_path))

    # Validation
    if not validation.empty:
        v = validation.dropna(subset=["Published", "Calculated"])
        if not v.empty:
            val_path = out_dir / f"{stem}_Validation.png"
            x = np.arange(len(v))
            width = 0.36
            fig, ax = plt.subplots(figsize=(max(9.5, len(v)*1.5), 5.8))
            ax.bar(x - width/2, v["Published"], width, label="Published")
            ax.bar(x + width/2, v["Calculated"], width, label="Calculated")
            ax.set_xticks(x)
            ax.set_xticklabels(v["Metric"], rotation=30, ha="right")
            ax.set_title("Published benchmark vs Universal TEA V0.6")
            ax.legend()
            fig.tight_layout()
            fig.savefig(val_path, dpi=220)
            plt.close(fig)
            paths.append(str(val_path))

    return paths


# ---------------------------------------------------------------------
# Results Excel
# ---------------------------------------------------------------------

def format_results_workbook(path: Path):
    wb = load_workbook(path)
    navy = "17365D"
    light = "D9EAF7"
    green = "E2F0D9"
    thin = Side(style="thin", color="D9D9D9")

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        # header row
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.freeze_panes = "A2"
        # widths
        for col in range(1, ws.max_column + 1):
            max_len = 0
            for row in range(1, min(ws.max_row, 80) + 1):
                val = ws.cell(row, col).value
                if val is not None:
                    max_len = max(max_len, len(str(val)))
            ws.column_dimensions[get_column_letter(col)].width = min(max(max_len + 2, 12), 38)
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="center")
                cell.border = Border(bottom=thin)
                if isinstance(cell.value, (int, float)):
                    # financial model convention: zeros as dash; negatives red/parentheses
                    cell.number_format = '#,##0.00;[Red](#,##0.00);-'
    wb.save(path)


def write_results_excel(
    path: Path,
    ctx: ModelContext,
    baseline: Dict[str, object],
    sensitivity: pd.DataFrame,
    scenarios: pd.DataFrame,
    break_even: pd.DataFrame,
    scale: pd.DataFrame,
    environmental: pd.DataFrame,
    validation: pd.DataFrame,
):
    summary = pd.DataFrame([
        ["Case", ctx.case_name, "-"],
        ["Process category", ctx.process_category, "-"],
        ["Primary product", ctx.primary_product, "-"],
        ["Operating hours", ctx.hours, "h/year"],
        ["Project life", ctx.life, "years"],
        ["Purchased equipment / base capital", ctx.purchased_equipment, ctx.currency],
        ["Installed equipment / cost blocks", ctx.installed_equipment, ctx.currency],
        ["FCI", ctx.fci, ctx.currency],
        ["TCI", baseline["TCI"], ctx.currency],
        ["Annual OPEX", baseline["Annual OPEX"], f"{ctx.currency}/year"],
        ["Annual revenue", baseline["Annual Revenue"], f"{ctx.currency}/year"],
        ["NPV", baseline["NPV"], ctx.currency],
        ["IRR", baseline["IRR"], "fraction"],
        ["Payback", baseline["Payback"], "years"],
        ["MSP", baseline["MSP"], f"{ctx.currency}/product unit"],
        ["Entered selling price", ctx.selling_price, ctx.selling_price_unit],
        ["Benefit-cost ratio", baseline["BCR"], "ratio"],
        ["Finance mode", ctx.finance_mode, "-"],
    ], columns=["Metric", "Value", "Unit"])

    opex_df = pd.DataFrame(
        [{"OPEX Category": k, "Annual Cost": v} for k, v in baseline["OPEX Components"].items()]
    )

    warnings_df = pd.DataFrame({"Warning": ctx.warnings}) if ctx.warnings else pd.DataFrame({"Warning":["None"]})

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        ctx.equipment_df.to_excel(writer, sheet_name="Equipment_Costs", index=False)
        opex_df.to_excel(writer, sheet_name="OPEX_Breakdown", index=False)
        baseline["Cash Flow"].to_excel(writer, sheet_name="Cash_Flow", index=False)
        sensitivity.to_excel(writer, sheet_name="Sensitivity", index=False)
        scenarios.to_excel(writer, sheet_name="Scenarios", index=False)
        break_even.to_excel(writer, sheet_name="Break_Even", index=False)
        scale.to_excel(writer, sheet_name="Scale_Analysis", index=False)
        environmental.to_excel(writer, sheet_name="Environmental", index=False)
        validation.to_excel(writer, sheet_name="Validation_Comparison", index=False)
        warnings_df.to_excel(writer, sheet_name="Warnings", index=False)

    format_results_workbook(path)


# ---------------------------------------------------------------------
# Word report
# ---------------------------------------------------------------------

def add_doc_table(doc, headers, rows, font_size=7.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = str(h)
        for run in table.rows[0].cells[i].paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(font_size)
    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            if isinstance(v, (float, np.floating)):
                if np.isnan(v):
                    text = "-"
                elif abs(v) >= 1e6:
                    text = f"{v:,.0f}"
                elif abs(v) >= 100:
                    text = f"{v:,.2f}"
                else:
                    text = f"{v:.4f}"
            else:
                text = "-" if v is None else str(v)
            cells[i].text = text
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(font_size)
    return table


def generate_verdict(
    ctx: ModelContext,
    baseline: Dict[str, object],
    scenarios: pd.DataFrame,
    sensitivity: pd.DataFrame,
    validation: pd.DataFrame,
) -> str:
    if not validation.empty:
        msp_rows = validation[
            validation["Metric"].str.lower().str.contains("production cost|msp", regex=True)
        ]
        if not msp_rows.empty and not np.isnan(msp_rows.iloc[0]["Deviation (%)"]):
            d = float(msp_rows.iloc[0]["Deviation (%)"])
            if abs(d) <= 1:
                quality = "excellent external agreement"
            elif abs(d) <= 5:
                quality = "good screening-level external agreement"
            else:
                quality = "a material external-validation gap"
            text = (
                f"The workbook-driven V0.6 model gives an MSP of {baseline['MSP']:.4f} "
                f"{ctx.currency}/product unit. Against the published benchmark, the primary cost deviation is "
                f"{d:+.2f}%, which represents {quality}. "
            )
        else:
            text = "Published validation targets are present, but the primary MSP target could not be mapped automatically. "

        if ctx.warnings:
            text += "Important V0.6 limitations: " + " ".join(ctx.warnings)
        else:
            text += "No unsupported finance/degradation inputs were detected for this case."
        return text

    viable = baseline["NPV"] >= 0 and (np.isnan(baseline["IRR"]) or baseline["IRR"] >= ctx.target_return)
    scenario_viable = 0
    scenario_total = 0
    if not scenarios.empty:
        scenario_total = len(scenarios)
        scenario_viable = int((scenarios["Status"] == "Viable").sum())
    top = sensitivity.iloc[0]["Parameter"] if not sensitivity.empty else "not evaluated"
    return (
        f"The base case is {'economically viable' if viable else 'economically weak'} at the entered selling price. "
        f"NPV = {baseline['NPV']:,.0f} {ctx.currency}, IRR = "
        f"{baseline['IRR']:.2%} and MSP = {baseline['MSP']:.4f} {ctx.currency}/product unit. "
        f"{scenario_viable} of {scenario_total} tested scenarios are viable. "
        f"The highest screened MSP influence is {top}. "
        "Environmental outputs, when enabled, are screening indicators rather than a full LCA."
    )


def write_word_report(
    path: Path,
    ctx: ModelContext,
    baseline: Dict[str, object],
    sensitivity: pd.DataFrame,
    scenarios: pd.DataFrame,
    break_even: pd.DataFrame,
    environmental: pd.DataFrame,
    validation: pd.DataFrame,
):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.55)
    sec.bottom_margin = Inches(0.55)
    sec.left_margin = Inches(0.55)
    sec.right_margin = Inches(0.55)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(9.2)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.12

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Universal TEA V0.6 — Decision and Validation Report")
    r.bold = True
    r.font.size = Pt(15)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(ctx.case_name)
    r.italic = True
    r.font.size = Pt(9)

    # Table 1 — Inputs
    p = doc.add_paragraph()
    rr = p.add_run("Table 1. Key input assumptions")
    rr.bold = True
    rr.font.size = Pt(11)
    inputs = [
        ["Process category", ctx.process_category, "-"],
        ["Operating hours", ctx.hours, "h/year"],
        ["Project life", ctx.life, "years"],
        ["Primary product", ctx.primary_product, "-"],
        ["Product rate", ctx.product_rate, ctx.product_unit],
        ["Selling price", ctx.selling_price, ctx.selling_price_unit],
        ["Target return", ctx.target_return, "fraction"],
        ["Tax rate", ctx.tax_rate, "fraction"],
        ["Finance mode", ctx.finance_mode, "-"],
        ["Installed capital", ctx.installed_equipment, ctx.currency],
    ]
    add_doc_table(doc, ["Input", "Value", "Unit"], inputs, 7.6)

    # Table 2 — Outputs
    p = doc.add_paragraph()
    rr = p.add_run("Table 2. Main TEA outputs")
    rr.bold = True
    rr.font.size = Pt(11)
    outputs = [
        ["TCI", baseline["TCI"], ctx.currency],
        ["Annual OPEX", baseline["Annual OPEX"], f"{ctx.currency}/year"],
        ["Annual revenue", baseline["Annual Revenue"], f"{ctx.currency}/year"],
        ["NPV", baseline["NPV"], ctx.currency],
        ["IRR", baseline["IRR"], "fraction"],
        ["Payback", baseline["Payback"], "years"],
        ["MSP", baseline["MSP"], f"{ctx.currency}/product unit"],
        ["Benefit-cost ratio", baseline["BCR"], "ratio"],
    ]
    if not environmental.empty:
        for name in ["Carbon intensity", "Carbon saving vs reference", "Resource efficiency"]:
            row = environmental[environmental["Indicator"].eq(name)]
            if not row.empty:
                outputs.append([name, row.iloc[0]["Value"], row.iloc[0]["Unit"]])
    add_doc_table(doc, ["Output", "Value", "Unit"], outputs, 7.6)

    doc.add_page_break()

    # Table 3 — Decision/validation index
    p = doc.add_paragraph()
    rr = p.add_run("Table 3. Decision-analysis / validation index")
    rr.bold = True
    rr.font.size = Pt(11)

    decision_rows = []
    if not validation.empty:
        for _, row in validation.iterrows():
            decision_rows.append([
                row["Metric"], row["Published"], row["Calculated"],
                row["Deviation (%)"], row["Status"]
            ])
        headers = ["Metric", "Published", "Calculated", "Deviation (%)", "Status"]
    else:
        if not sensitivity.empty:
            for _, row in sensitivity.head(3).iterrows():
                decision_rows.append([
                    f"Sensitivity rank {int(row['Rank'])}", row["Parameter"],
                    row["Max |MSP Change|"], row["Priority"]
                ])
        if not break_even.empty:
            for _, row in break_even.head(3).iterrows():
                decision_rows.append([
                    "Break-even", row["Parameter"], row["Break-even Factor"], row["Status"]
                ])
        if not scenarios.empty:
            decision_rows.append([
                "Scenario viability", f"{int((scenarios['Status']=='Viable').sum())}/{len(scenarios)}",
                "viable scenarios", "-"
            ])
        headers = ["Index", "Item", "Result", "Status"]
    add_doc_table(doc, headers, decision_rows, 7.1)

    p = doc.add_paragraph()
    rr = p.add_run("Final verdict")
    rr.bold = True
    rr.font.size = Pt(11)
    doc.add_paragraph(generate_verdict(ctx, baseline, scenarios, sensitivity, validation))

    if ctx.warnings:
        p = doc.add_paragraph()
        rr = p.add_run("Methodology warnings")
        rr.bold = True
        rr.font.size = Pt(10)
        for warning in ctx.warnings:
            doc.add_paragraph(warning, style=None)

    doc.save(path)


# ---------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="Universal TEA Engine V0.6 — workbook-driven TEA and validation"
    )
    parser.add_argument("input_workbook", help="Universal TEA V0.6 input .xlsx")
    parser.add_argument("--output-dir", default=".", help="Output directory")
    parser.add_argument("--no-report", action="store_true", help="Do not generate Word report")
    parser.add_argument("--no-figures", action="store_true", help="Do not generate PNG figures")
    args = parser.parse_args()

    input_path = Path(args.input_workbook)
    if not input_path.exists():
        raise SystemExit(f"Input workbook not found: {input_path}")

    out_dir = Path(args.output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    ctx, rd, decision_df, validation_targets = parse_model(str(input_path))
    stem = slugify(input_path.stem)

    baseline = evaluate_case(ctx)
    sensitivity = run_sensitivity(ctx, decision_df)
    scenarios = run_scenarios(ctx, decision_df)
    break_even = run_break_even(ctx, decision_df)
    scale = run_scale(ctx, decision_df)
    environmental = calculate_environmental(ctx, rd)
    validation = validation_comparison(ctx, validation_targets, baseline, sensitivity)

    results_path = out_dir / f"{stem}_TEA_Results_V0_6.xlsx"
    report_path = out_dir / f"{stem}_TEA_Report_V0_6.docx"

    write_results_excel(
        results_path, ctx, baseline, sensitivity, scenarios,
        break_even, scale, environmental, validation
    )

    if not args.no_report:
        write_word_report(
            report_path, ctx, baseline, sensitivity, scenarios,
            break_even, environmental, validation
        )

    figure_paths = []
    if not args.no_figures:
        figure_paths = save_charts(
            ctx, out_dir, stem, baseline, sensitivity, scenarios, validation
        )

    print("=" * 72)
    print("UNIVERSAL TEA ENGINE V0.6")
    print("=" * 72)
    print(f"Case                 : {ctx.case_name}")
    print(f"Finance mode         : {ctx.finance_mode}")
    print(f"TCI                  : {baseline['TCI']:,.2f} {ctx.currency}")
    print(f"Annual OPEX          : {baseline['Annual OPEX']:,.2f} {ctx.currency}/year")
    print(f"Annual revenue       : {baseline['Annual Revenue']:,.2f} {ctx.currency}/year")
    print(f"NPV                  : {baseline['NPV']:,.2f} {ctx.currency}")
    print(f"IRR                  : {baseline['IRR']:.4%}" if not np.isnan(baseline["IRR"]) else "IRR                  : n/a")
    print(f"Payback              : {baseline['Payback']:.3f} years" if not np.isnan(baseline["Payback"]) else "Payback              : not reached")
    print(f"MSP                  : {baseline['MSP']:.6f} {ctx.currency}/product unit")
    print(f"Entered selling price: {ctx.selling_price:.6f} {ctx.selling_price_unit}")
    if not validation.empty:
        print("\nValidation comparison:")
        print(validation[["Metric","Published","Calculated","Deviation (%)","Status"]].to_string(index=False))
    if ctx.warnings:
        print("\nWarnings:")
        for w in ctx.warnings:
            print(" -", w)
    print("\nOutputs:")
    print(f" - {results_path}")
    if not args.no_report:
        print(f" - {report_path}")
    for p in figure_paths:
        print(f" - {p}")




# =====================================================================
# V0.6 DECISION-SUPPORT EXTENSION
# =====================================================================
# V0.5e deterministic TEA equations are retained. V0.6 adds:
# - enhanced technically coupled sensitivity
# - Monte Carlo uncertainty
# - probabilistic viability
# - uncertainty-driver correlations
# - risk and R&D priority indices
# - carbon-value screening
# - consolidated decision summary
# - journal-standard figures
#
# IMPORTANT:
# * Monte Carlo inputs are sampled independently in V0.6.
# * Risk Priority Score and R&D Priority Score are transparent,
#   software-defined decision-support indices, not industry standards.
# * Environmental calculations remain screening-level and are not a full LCA.
# =====================================================================

import textwrap as _textwrap


# Keep access to the V0.5e environmental routine for extension.
_calculate_environmental_v05e = calculate_environmental


def _norm_parameter_name(value: str) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def parameter_mods(ctx: ModelContext, parameter: str, factor: float) -> Dict[str, float]:
    """
    V0.6 technically coupled parameter mapping.

    Capacity/availability:
        scales annual operating hours, product output and flow-based variable OPEX
        through the existing hours_factor relationship; fixed OPEX is unchanged.

    Product output/yield:
        changes product output while keeping feed/utility quantities at base values,
        representing yield/conversion uncertainty.

    Process efficiency:
        for generic active-flow cases, specific material and utility demand change
        inversely with the efficiency factor. Published H2A biomass cases retain the
        feedstock-cost decomposition logic.
    """
    p = _norm_parameter_name(parameter)
    mods: Dict[str, float] = {}

    if p in {"equipment capex", "total capital investment"}:
        mods["tci_factor"] = factor
    elif p in {"material prices", "feedstock price"}:
        mods["material_factor"] = factor
    elif p == "utility prices":
        mods["utility_factor"] = factor
    elif p == "product selling price":
        mods["price"] = ctx.selling_price * factor
    elif p in {"operating hours", "operating capacity factor", "capacity factor"}:
        mods["hours_factor"] = factor
    elif p in {"primary product output", "product output", "product yield", "conversion yield", "yield"}:
        mods["product_output_factor"] = factor
    elif p == "electricity demand":
        mods["electricity_demand_factor"] = factor
    elif p in {"fixed opex", "total fixed operating cost"}:
        mods["fixed_factor"] = factor

    elif p == "biomass feedstock price":
        ref = find_reference_cost(ctx.reference_material_costs, "biomass")
        if ref > 0 and ctx.direct_variable_opex > 0:
            mods["variable_opex_adjustment"] = ref * (factor - 1.0)
        else:
            mods["material_factor"] = factor

    elif p in {"plant efficiency", "process efficiency", "conversion efficiency"}:
        ref = find_reference_cost(ctx.reference_material_costs, "biomass")
        if ref > 0 and ctx.direct_variable_opex > 0:
            if factor > 0:
                mods["variable_opex_adjustment"] = ref * ((1.0 / factor) - 1.0)
        elif factor > 0:
            # Generic coupling: improved efficiency reduces specific material and
            # utility requirements for the same nominal product rate.
            inv = 1.0 / factor
            mods["material_factor"] = inv
            mods["utility_factor"] = inv
            mods["electricity_demand_factor"] = inv

    elif p == "labor requirement":
        mods["fixed_factor"] = factor

    elif p == "electricity price":
        ref = find_reference_cost(ctx.reference_utility_costs, "electric")
        if ref > 0 and ctx.direct_variable_opex > 0:
            mods["variable_opex_adjustment"] = ref * (factor - 1.0)
        else:
            mods["electricity_factor"] = factor

    elif p == "stack electrical usage":
        ref = find_reference_cost(ctx.reference_utility_costs, "electric")
        if ref > 0 and ctx.direct_variable_opex > 0:
            mods["variable_opex_adjustment"] = ref * (factor - 1.0)
        else:
            mods["electricity_demand_factor"] = factor

    elif p == "stack cost":
        stack_installed = find_reference_cost(ctx.reference_equipment_installed, "stack")
        mods["tci_adjustment"] = stack_installed * (factor - 1.0)

    elif p == "stack replacement cost":
        mods["replacement_fraction_factor"] = factor

    elif p == "stack replacement interval":
        mods["replacement_interval_factor"] = factor

    return mods


def run_sensitivity(ctx: ModelContext, decision_df: pd.DataFrame) -> pd.DataFrame:
    """
    Enhanced V0.6 sensitivity:
    - low/high NPV and MSP
    - NPV impact normalized by TCI
    - MSP impact normalized by baseline MSP
    - combined normalized impact score
    """
    if decision_df.empty:
        return pd.DataFrame()
    rows = decision_df[decision_df["Section"].eq("Sensitivity")] if "Section" in decision_df.columns else pd.DataFrame()
    if rows.empty:
        return pd.DataFrame()

    out = []
    baseline = evaluate_case(ctx)
    tci_scale = max(abs(float(baseline["TCI"])), 1.0)
    msp_scale = max(abs(float(baseline["MSP"])), 1e-9)

    for _, row in rows.iterrows():
        if not yes(row.get("Analyze?")):
            continue
        parameter = str(row.get("Parameter") or "")
        low_factor = fnum(row.get("Low Factor"), 1.0)
        high_factor = fnum(row.get("High Factor"), 1.0)

        low = evaluate_case(ctx, **parameter_mods(ctx, parameter, low_factor))
        high = evaluate_case(ctx, **parameter_mods(ctx, parameter, high_factor))

        max_npv_change = max(
            abs(float(low["NPV"]) - float(baseline["NPV"])),
            abs(float(high["NPV"]) - float(baseline["NPV"])),
        )
        max_msp_change = max(
            abs(float(low["MSP"]) - float(baseline["MSP"])),
            abs(float(high["MSP"]) - float(baseline["MSP"])),
        )

        npv_rel = max_npv_change / tci_scale
        msp_rel = max_msp_change / msp_scale
        raw = max(npv_rel, msp_rel)

        out.append({
            "Parameter": parameter,
            "Low Factor": low_factor,
            "High Factor": high_factor,
            "Low NPV": low["NPV"],
            "Baseline NPV": baseline["NPV"],
            "High NPV": high["NPV"],
            "Low MSP": low["MSP"],
            "Baseline MSP": baseline["MSP"],
            "High MSP": high["MSP"],
            "Low MSP Delta": low["MSP"] - baseline["MSP"],
            "High MSP Delta": high["MSP"] - baseline["MSP"],
            "Max |NPV Change|": max_npv_change,
            "Max |MSP Change|": max_msp_change,
            "NPV Impact / TCI": npv_rel,
            "MSP Relative Impact": msp_rel,
            "Combined Impact Raw": raw,
        })

    if not out:
        return pd.DataFrame()

    df = pd.DataFrame(out).sort_values("Combined Impact Raw", ascending=False).reset_index(drop=True)
    max_raw = max(float(df["Combined Impact Raw"].max()), 1e-12)
    df["Impact Score"] = df["Combined Impact Raw"] / max_raw
    df["Priority"] = df["Impact Score"].apply(
        lambda x: "High" if x >= 0.67 else ("Medium" if x >= 0.33 else "Low")
    )
    df["Rank"] = np.arange(1, len(df) + 1)
    return df


def parse_uncertainty_settings(rd: WorkbookReader) -> Dict[str, object]:
    rows = rd.rows_by_header("Uncertainty_Settings")
    settings: Dict[str, object] = {}
    for row in rows:
        key = row.get("Parameter")
        if key not in (None, ""):
            settings[_norm_parameter_name(key)] = row.get("Value")
    return settings


def parse_uncertainty_parameters(rd: WorkbookReader) -> pd.DataFrame:
    rows = rd.rows_by_header("Uncertainty_Parameters")
    if not rows:
        return pd.DataFrame()
    return pd.DataFrame(rows)


def _merge_mods(base: Dict[str, float], new: Dict[str, float]) -> Dict[str, float]:
    factor_keys = {
        "tci_factor","hours_factor","product_output_factor","material_factor",
        "utility_factor","electricity_factor","electricity_demand_factor",
        "fixed_factor","direct_variable_factor","replacement_fraction_factor",
        "replacement_interval_factor"
    }
    additive_keys = {"tci_adjustment","variable_opex_adjustment"}

    out = dict(base)
    for key, value in new.items():
        if key in factor_keys:
            out[key] = out.get(key, 1.0) * float(value)
        elif key in additive_keys:
            out[key] = out.get(key, 0.0) + float(value)
        elif key == "price":
            out[key] = float(value)
        else:
            out[key] = value
    return out


def _sample_factor(rng, distribution: str, low: float, mode: float, high: float, n: int) -> np.ndarray:
    d = _norm_parameter_name(distribution)
    low, mode, high = float(low), float(mode), float(high)
    if high < low:
        low, high = high, low
    mode = min(max(mode, low), high)

    if d in {"triangular", "triangle"}:
        return rng.triangular(low, mode, high, size=n)
    if d == "uniform":
        return rng.uniform(low, high, size=n)
    if d in {"normal", "gaussian"}:
        sigma = (high - low) / 6.0 if high > low else 0.0
        if sigma <= 0:
            return np.full(n, mode)
        return np.clip(rng.normal(mode, sigma, size=n), low, high)
    return np.full(n, mode)


def evaluate_case_fast(ctx: ModelContext, **mods) -> Dict[str, float]:
    mods = dict(mods)
    price = mods.pop("price", ctx.selling_price)

    cashflows = fast_cashflow_values(ctx, price, **mods)
    model_npv = npv(cashflows, ctx.target_return)
    irr = solve_irr(cashflows)

    def builder(p):
        return fast_cashflow_values(ctx, p, **mods)

    msp = solve_root_price(builder, ctx.target_return)

    annual_product = annual_product_quantity(
        ctx,
        mods.get("hours_factor", 1.0),
        mods.get("product_output_factor", 1.0),
    )
    comps = opex_components(
        ctx,
        hours_factor=mods.get("hours_factor", 1.0),
        product_output_factor=mods.get("product_output_factor", 1.0),
        material_factor=mods.get("material_factor", 1.0),
        utility_factor=mods.get("utility_factor", 1.0),
        electricity_factor=mods.get("electricity_factor", 1.0),
        electricity_demand_factor=mods.get("electricity_demand_factor", 1.0),
        fixed_factor=mods.get("fixed_factor", 1.0),
        direct_variable_factor=mods.get("direct_variable_factor", 1.0),
    )
    comps["Direct variable OPEX"] += mods.get("variable_opex_adjustment", 0.0)
    annual_opex = total_opex(comps)
    annual_revenue = annual_product * price
    tci = ctx.tci * mods.get("tci_factor", 1.0) + mods.get("tci_adjustment", 0.0)

    return {
        "NPV": model_npv,
        "IRR": irr,
        "MSP": msp,
        "Annual Product": annual_product,
        "Annual OPEX": annual_opex,
        "Annual Revenue": annual_revenue,
        "TCI": tci,
        "Price": price,
    }


def run_uncertainty(
    ctx: ModelContext,
    rd: WorkbookReader,
    baseline: Dict[str, object],
) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    settings = parse_uncertainty_settings(rd)
    params = parse_uncertainty_parameters(rd)

    if _norm_parameter_name(settings.get("uncertainty active", "No")) not in {"yes","y","true","1","active"}:
        return pd.DataFrame(), pd.DataFrame(), pd.DataFrame()
    if params.empty:
        return pd.DataFrame(), pd.DataFrame(), pd.DataFrame()

    active = params[params["Active?"].apply(yes)].copy()
    if active.empty:
        return pd.DataFrame(), pd.DataFrame(), pd.DataFrame()

    n = int(max(100, fnum(settings.get("simulation count"), 1000)))
    seed = int(fnum(settings.get("random seed"), 42))
    p10 = fnum(settings.get("p10 percentile"), 10.0)
    p50 = fnum(settings.get("p50 percentile"), 50.0)
    p90 = fnum(settings.get("p90 percentile"), 90.0)

    rng = np.random.default_rng(seed)

    factor_arrays: Dict[str, np.ndarray] = {}
    for _, row in active.iterrows():
        name = str(row.get("Parameter") or "")
        factor_arrays[name] = _sample_factor(
            rng,
            str(row.get("Distribution") or "Triangular"),
            fnum(row.get("Low Factor"), 1.0),
            fnum(row.get("Most Likely Factor"), 1.0),
            fnum(row.get("High Factor"), 1.0),
            n
        )

    records = []
    for i in range(n):
        mods: Dict[str, float] = {}
        rec = {"Simulation": i + 1}
        for name, arr in factor_arrays.items():
            factor = float(arr[i])
            rec[f"Factor | {name}"] = factor
            mods = _merge_mods(mods, parameter_mods(ctx, name, factor))
        res = evaluate_case_fast(ctx, **mods)
        rec.update({
            "NPV": res["NPV"],
            "IRR": res["IRR"],
            "MSP": res["MSP"],
            "Annual Product": res["Annual Product"],
            "Annual OPEX": res["Annual OPEX"],
            "Annual Revenue": res["Annual Revenue"],
            "TCI": res["TCI"],
            "Sampled Selling Price": res["Price"],
        })
        records.append(rec)

    samples = pd.DataFrame(records)

    def q(series, p):
        vals = pd.to_numeric(series, errors="coerce").dropna()
        return float(np.percentile(vals, p)) if len(vals) else np.nan

    valid_irr = pd.to_numeric(samples["IRR"], errors="coerce")
    prob_npv = float((samples["NPV"] > 0).mean())
    prob_irr = float((valid_irr >= ctx.target_return).mean())
    prob_msp = float((samples["MSP"] <= ctx.selling_price).mean())

    summary_rows = [
        ["Simulation count", n, "runs"],
        ["Random seed", seed, "integer"],
        ["Sampling dependence", "Independent", "V0.6 scope"],
        ["Baseline NPV", baseline["NPV"], ctx.currency],
        ["Mean NPV", samples["NPV"].mean(), ctx.currency],
        [f"P{int(p10)} NPV", q(samples["NPV"], p10), ctx.currency],
        [f"P{int(p50)} NPV", q(samples["NPV"], p50), ctx.currency],
        [f"P{int(p90)} NPV", q(samples["NPV"], p90), ctx.currency],
        ["NPV standard deviation", samples["NPV"].std(ddof=1), ctx.currency],
        ["Baseline MSP", baseline["MSP"], f"{ctx.currency}/product unit"],
        ["Mean MSP", samples["MSP"].mean(), f"{ctx.currency}/product unit"],
        [f"P{int(p10)} MSP", q(samples["MSP"], p10), f"{ctx.currency}/product unit"],
        [f"P{int(p50)} MSP", q(samples["MSP"], p50), f"{ctx.currency}/product unit"],
        [f"P{int(p90)} MSP", q(samples["MSP"], p90), f"{ctx.currency}/product unit"],
        ["MSP standard deviation", samples["MSP"].std(ddof=1), f"{ctx.currency}/product unit"],
        ["Mean IRR", valid_irr.mean(), "fraction"],
        [f"P{int(p10)} IRR", q(valid_irr, p10), "fraction"],
        [f"P{int(p50)} IRR", q(valid_irr, p50), "fraction"],
        [f"P{int(p90)} IRR", q(valid_irr, p90), "fraction"],
        ["P(NPV > 0)", prob_npv, "probability"],
        ["P(IRR >= target return)", prob_irr, "probability"],
        ["P(MSP <= entered selling price)", prob_msp, "probability"],
    ]
    summary = pd.DataFrame(summary_rows, columns=["Indicator","Value","Unit"])

    # Spearman rank correlations provide an uncertainty-driver diagnostic.
    driver_rows = []
    for name in factor_arrays:
        col = f"Factor | {name}"
        rho_npv = samples[[col, "NPV"]].corr(method="spearman").iloc[0,1]
        rho_msp = samples[[col, "MSP"]].corr(method="spearman").iloc[0,1]
        driver_rows.append({
            "Parameter": name,
            "Spearman rho vs NPV": rho_npv,
            "|rho| vs NPV": abs(rho_npv) if not np.isnan(rho_npv) else np.nan,
            "Spearman rho vs MSP": rho_msp,
            "|rho| vs MSP": abs(rho_msp) if not np.isnan(rho_msp) else np.nan,
        })
    drivers = pd.DataFrame(driver_rows)
    if not drivers.empty:
        drivers["Max |rho|"] = drivers[["|rho| vs NPV","|rho| vs MSP"]].max(axis=1)
        drivers["Dominant Outcome"] = np.where(
            drivers["|rho| vs NPV"] >= drivers["|rho| vs MSP"], "NPV", "MSP"
        )
        drivers = drivers.sort_values("Max |rho|", ascending=False).reset_index(drop=True)
        drivers["Driver Rank"] = np.arange(1, len(drivers)+1)

    return summary, samples, drivers


def build_rd_priority(
    ctx: ModelContext,
    baseline: Dict[str, object],
    sensitivity: pd.DataFrame,
    rd: WorkbookReader,
    drivers: pd.DataFrame,
) -> pd.DataFrame:
    params = parse_uncertainty_parameters(rd)
    if params.empty:
        return pd.DataFrame()
    active = params[params["Active?"].apply(yes)].copy()
    if active.empty:
        return pd.DataFrame()

    raw_rows = []
    for _, row in active.iterrows():
        parameter = str(row.get("Parameter") or "")
        low = fnum(row.get("Low Factor"), 1.0)
        mode = fnum(row.get("Most Likely Factor"), 1.0)
        high = fnum(row.get("High Factor"), 1.0)
        width = abs(high - low) / max(abs(mode), 1e-9)

        match = sensitivity[
            sensitivity["Parameter"].apply(_norm_parameter_name).eq(_norm_parameter_name(parameter))
        ] if not sensitivity.empty else pd.DataFrame()

        if not match.empty:
            impact_raw = float(match.iloc[0]["Combined Impact Raw"])
        else:
            low_res = evaluate_case_fast(ctx, **parameter_mods(ctx, parameter, low))
            high_res = evaluate_case_fast(ctx, **parameter_mods(ctx, parameter, high))
            tci_scale = max(abs(float(baseline["TCI"])), 1.0)
            msp_scale = max(abs(float(baseline["MSP"])), 1e-9)
            npv_rel = max(
                abs(low_res["NPV"] - baseline["NPV"]),
                abs(high_res["NPV"] - baseline["NPV"])
            ) / tci_scale
            msp_rel = max(
                abs(low_res["MSP"] - baseline["MSP"]),
                abs(high_res["MSP"] - baseline["MSP"])
            ) / msp_scale
            impact_raw = max(npv_rel, msp_rel)

        rho_npv = np.nan
        rho_msp = np.nan
        if not drivers.empty:
            dm = drivers[drivers["Parameter"].apply(_norm_parameter_name).eq(_norm_parameter_name(parameter))]
            if not dm.empty:
                rho_npv = float(dm.iloc[0]["Spearman rho vs NPV"])
                rho_msp = float(dm.iloc[0]["Spearman rho vs MSP"])

        controllability = str(row.get("Controllability") or "Mixed")
        rd_weight = fnum(row.get("R&D Weight"), 0.6)

        raw_rows.append({
            "Parameter": parameter,
            "Category": row.get("Category"),
            "Distribution": row.get("Distribution"),
            "Low Factor": low,
            "Most Likely Factor": mode,
            "High Factor": high,
            "Controllability": controllability,
            "R&D Weight": rd_weight,
            "Impact Raw": impact_raw,
            "Uncertainty Width Raw": width,
            "Spearman rho vs NPV": rho_npv,
            "Spearman rho vs MSP": rho_msp,
        })

    df = pd.DataFrame(raw_rows)
    max_imp = max(float(df["Impact Raw"].max()), 1e-12)
    max_width = max(float(df["Uncertainty Width Raw"].max()), 1e-12)
    df["Sensitivity Impact Score"] = df["Impact Raw"] / max_imp
    df["Uncertainty Width Score"] = df["Uncertainty Width Raw"] / max_width
    df["Risk Priority Raw"] = df["Sensitivity Impact Score"] * df["Uncertainty Width Score"]
    max_risk = max(float(df["Risk Priority Raw"].max()), 1e-12)
    df["Risk Priority Score"] = df["Risk Priority Raw"] / max_risk
    df["R&D Priority Raw"] = df["Risk Priority Raw"] * df["R&D Weight"]
    max_rd = max(float(df["R&D Priority Raw"].max()), 1e-12)
    df["R&D Priority Score"] = df["R&D Priority Raw"] / max_rd

    settings = parse_uncertainty_settings(rd)
    high_t = fnum(settings.get("r&d priority high threshold"), 0.67)
    medium_t = fnum(settings.get("r&d priority medium threshold"), 0.33)
    df["R&D Priority"] = df["R&D Priority Score"].apply(
        lambda x: "High" if x >= high_t else ("Medium" if x >= medium_t else "Low")
    )
    df = df.sort_values("R&D Priority Score", ascending=False).reset_index(drop=True)
    df["R&D Rank"] = np.arange(1, len(df)+1)

    def action(row):
        c = _norm_parameter_name(row["Controllability"])
        if "research" in c:
            return "Prioritize experiments, process optimization and validation."
        if "external" in c or "market" in c:
            return "Treat mainly as market/procurement risk; monitor, contract or hedge rather than relying on R&D."
        return "Use combined R&D, design optimization and operational/procurement risk management."

    df["Recommended Action"] = df.apply(action, axis=1)
    return df


def calculate_environmental(
    ctx: ModelContext,
    rd: WorkbookReader,
    baseline: Optional[Dict[str, object]] = None,
) -> pd.DataFrame:
    base = _calculate_environmental_v05e(ctx, rd)
    if base.empty:
        return base

    rows = rd.rows_by_header("Environmental")
    carbon_price = 0.0
    carbon_price_active = False
    for row in rows:
        if _norm_parameter_name(row.get("Parameter")) == "carbon price":
            carbon_price_active = yes(row.get("Active?"))
            carbon_price = fnum(row.get("Value"), 0.0)
            break

    if carbon_price_active and carbon_price != 0:
        saving_row = base[base["Indicator"].eq("Annual carbon saving")]
        if not saving_row.empty:
            annual_saving_kg = fnum(saving_row.iloc[0]["Value"], np.nan)
            if not np.isnan(annual_saving_kg):
                annual_value = annual_saving_kg / 1000.0 * carbon_price
                pv_value = sum(
                    annual_value / ((1 + ctx.target_return) ** y)
                    for y in range(1, ctx.life + 1)
                )
                extra = [
                    ["Carbon price", carbon_price, f"{ctx.currency}/tCO2e"],
                    ["Annual carbon value", annual_value, f"{ctx.currency}/year"],
                    ["PV carbon benefit", pv_value, ctx.currency],
                ]
                if baseline is not None:
                    extra.append([
                        "Carbon-value-adjusted NPV",
                        float(baseline["NPV"]) + pv_value,
                        ctx.currency
                    ])
                base = pd.concat(
                    [base, pd.DataFrame(extra, columns=["Indicator","Value","Unit"])],
                    ignore_index=True
                )
    return base


def _get_indicator(df: pd.DataFrame, name: str, default=np.nan):
    if df is None or df.empty:
        return default
    row = df[df.iloc[:,0].astype(str).str.lower().eq(name.lower())]
    if row.empty:
        return default
    return row.iloc[0,1]


def build_decision_summary(
    ctx: ModelContext,
    baseline: Dict[str, object],
    sensitivity: pd.DataFrame,
    scenarios: pd.DataFrame,
    break_even: pd.DataFrame,
    uncertainty_summary: pd.DataFrame,
    rd_priority: pd.DataFrame,
    environmental: pd.DataFrame,
    validation: pd.DataFrame,
    rd: WorkbookReader,
) -> pd.DataFrame:
    base_viable = baseline["NPV"] >= 0 and (
        np.isnan(baseline["IRR"]) or baseline["IRR"] >= ctx.target_return
    )

    scenario_rate = np.nan
    if not scenarios.empty:
        scenario_rate = float((scenarios["Status"] == "Viable").mean())

    prob_npv = np.nan
    prob_irr = np.nan
    prob_msp = np.nan
    if not uncertainty_summary.empty:
        prob_npv = _get_indicator(uncertainty_summary, "P(NPV > 0)")
        prob_irr = _get_indicator(uncertainty_summary, "P(IRR >= target return)")
        prob_msp = _get_indicator(uncertainty_summary, "P(MSP <= entered selling price)")

    settings = parse_uncertainty_settings(rd)
    robust_t = fnum(settings.get("robust viability probability"), 0.80)
    conditional_t = fnum(settings.get("conditional viability probability"), 0.50)

    if not base_viable:
        robustness = "Economically weak"
    elif np.isnan(prob_npv):
        robustness = "Deterministically viable; uncertainty not enabled"
    elif prob_npv >= robust_t and (np.isnan(scenario_rate) or scenario_rate >= 0.50):
        robustness = "Robustly viable"
    elif prob_npv >= conditional_t:
        robustness = "Conditionally viable"
    else:
        robustness = "Viable base case but high downside risk"

    top_sens = sensitivity.iloc[0]["Parameter"] if not sensitivity.empty else "Not evaluated"
    top_rd = rd_priority.iloc[0]["Parameter"] if not rd_priority.empty else "Not evaluated"
    top_risk = (
        rd_priority.sort_values("Risk Priority Score", ascending=False).iloc[0]["Parameter"]
        if not rd_priority.empty else "Not evaluated"
    )

    closest_be = "Not bracketed"
    if not break_even.empty:
        solved = break_even[break_even["Status"].eq("Solved")].copy()
        if not solved.empty:
            solved["Abs Margin"] = solved["Change vs Base (%)"].abs()
            r = solved.sort_values("Abs Margin").iloc[0]
            closest_be = f"{r['Parameter']} ({r['Change vs Base (%)']:+.1f}% from base)"

    carbon_ci = _get_indicator(environmental, "Carbon intensity")
    carbon_save = _get_indicator(environmental, "Carbon saving vs reference")
    resource_eff = _get_indicator(environmental, "Resource efficiency")
    carbon_adj_npv = _get_indicator(environmental, "Carbon-value-adjusted NPV")

    validation_dev = np.nan
    if not validation.empty:
        m = validation[
            validation["Metric"].astype(str).str.lower().str.contains("production cost|msp", regex=True)
        ]
        if not m.empty:
            validation_dev = m.iloc[0]["Deviation (%)"]

    rows = [
        ["Base economic status", "Viable" if base_viable else "Economically weak", "-", "Economic"],
        ["Economic robustness class", robustness, "software-defined", "Decision"],
        ["Base NPV", baseline["NPV"], ctx.currency, "Economic"],
        ["Base IRR", baseline["IRR"], "fraction", "Economic"],
        ["Base MSP", baseline["MSP"], f"{ctx.currency}/product unit", "Economic"],
        ["Scenario viability rate", scenario_rate, "fraction", "Risk"],
        ["P(NPV > 0)", prob_npv, "probability", "Uncertainty"],
        ["P(IRR >= target return)", prob_irr, "probability", "Uncertainty"],
        ["P(MSP <= entered selling price)", prob_msp, "probability", "Uncertainty"],
        ["Top deterministic sensitivity", top_sens, "-", "Sensitivity"],
        ["Top uncertainty-risk parameter", top_risk, "-", "Risk"],
        ["Top R&D priority", top_rd, "-", "R&D"],
        ["Closest break-even margin", closest_be, "-", "Break-even"],
        ["Carbon intensity", carbon_ci, "kg CO2e/product unit", "Environmental"],
        ["Carbon saving vs reference", carbon_save, "%", "Environmental"],
        ["Resource efficiency", resource_eff, "fraction", "Environmental"],
        ["Carbon-value-adjusted NPV", carbon_adj_npv, ctx.currency, "Environmental"],
        ["Primary external-validation deviation", validation_dev, "%", "Validation"],
    ]
    return pd.DataFrame(rows, columns=["Decision Indicator","Value","Unit","Module"])


def _journal_label(value, width=25):
    return _textwrap.fill(str(value), width=width, break_long_words=False)


def _apply_journal_style(ax, title, xlabel=None, ylabel=None, grid_axis="x"):
    ax.set_title(title, fontsize=14, fontweight="semibold", pad=12)
    if xlabel:
        ax.set_xlabel(xlabel, fontsize=12, fontweight="semibold", labelpad=8)
    if ylabel:
        ax.set_ylabel(ylabel, fontsize=12, fontweight="semibold", labelpad=8)
    ax.tick_params(axis="both", labelsize=10.5)
    for lab in ax.get_xticklabels() + ax.get_yticklabels():
        lab.set_fontweight("medium")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.grid(True, axis=grid_axis, alpha=0.18, linewidth=0.8)
    ax.set_axisbelow(True)


def _save_figure(fig, path: Path):
    fig.savefig(path, dpi=400, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def save_charts(
    ctx: ModelContext,
    out_dir: Path,
    stem: str,
    baseline: Dict[str, object],
    sensitivity: pd.DataFrame,
    scenarios: pd.DataFrame,
    validation: pd.DataFrame,
    break_even: Optional[pd.DataFrame] = None,
    scale: Optional[pd.DataFrame] = None,
    environmental: Optional[pd.DataFrame] = None,
    uncertainty_summary: Optional[pd.DataFrame] = None,
    uncertainty_samples: Optional[pd.DataFrame] = None,
    uncertainty_drivers: Optional[pd.DataFrame] = None,
    rd_priority: Optional[pd.DataFrame] = None,
) -> List[str]:
    from matplotlib.ticker import FuncFormatter

    paths: List[str] = []

    def money_formatter(x, pos=None):
        a = abs(x)
        if a >= 1e9:
            return f"{x/1e9:.1f}B"
        if a >= 1e6:
            return f"{x/1e6:.1f}M"
        if a >= 1e3:
            return f"{x/1e3:.0f}k"
        return f"{x:.0f}"

    # CAPEX donut: legend outside, no wedge labels -> no overlap.
    if not ctx.equipment_df.empty and ctx.equipment_df["Installed Cost"].sum() > 0:
        data = ctx.equipment_df.groupby("Equipment")["Installed Cost"].sum().sort_values(ascending=False)
        if len(data) > 7:
            data = pd.concat([data.iloc[:7], pd.Series({"Other": data.iloc[7:].sum()})])
        total = data.sum()
        fig, ax = plt.subplots(figsize=(10.5, 6.8))
        wedges, _ = ax.pie(data.values, startangle=90, wedgeprops={"width": 0.38})
        legend_labels = [
            f"{_journal_label(k, 30)} — {v/total*100:.1f}%"
            for k, v in data.items()
        ]
        ax.legend(
            wedges, legend_labels, loc="center left", bbox_to_anchor=(1.02, 0.5),
            frameon=False, fontsize=10.5, handlelength=1.2
        )
        ax.text(0, 0.03, f"{ctx.currency}", ha="center", va="center",
                fontsize=11, fontweight="semibold")
        ax.text(0, -0.13, f"{total:,.0f}", ha="center", va="center",
                fontsize=13, fontweight="semibold")
        ax.set_title("Installed CAPEX breakdown", fontsize=14, fontweight="semibold", pad=12)
        fig.subplots_adjust(right=0.68)
        path = out_dir / f"{stem}_CAPEX.png"
        _save_figure(fig, path); paths.append(str(path))

    # OPEX donut.
    comps = pd.Series(baseline["OPEX Components"])
    comps = comps[comps > 0].sort_values(ascending=False)
    if len(comps):
        total = comps.sum()
        fig, ax = plt.subplots(figsize=(10.5, 6.8))
        wedges, _ = ax.pie(comps.values, startangle=90, wedgeprops={"width": 0.38})
        legend_labels = [
            f"{_journal_label(k, 30)} — {v/total*100:.1f}%"
            for k, v in comps.items()
        ]
        ax.legend(
            wedges, legend_labels, loc="center left", bbox_to_anchor=(1.02, 0.5),
            frameon=False, fontsize=10.5, handlelength=1.2
        )
        ax.text(0, 0.03, f"{ctx.currency}/y", ha="center", va="center",
                fontsize=11, fontweight="semibold")
        ax.text(0, -0.13, f"{total:,.0f}", ha="center", va="center",
                fontsize=13, fontweight="semibold")
        ax.set_title("Annual OPEX breakdown", fontsize=14, fontweight="semibold", pad=12)
        fig.subplots_adjust(right=0.68)
        path = out_dir / f"{stem}_OPEX.png"
        _save_figure(fig, path); paths.append(str(path))

    # Cash flow.
    cf = baseline["Cash Flow"]
    fig, ax = plt.subplots(figsize=(10.5, 6.2))
    ax.plot(cf["Time"], cf["Cumulative Cash Flow"], marker="o", markersize=3.2, linewidth=1.8)
    ax.axhline(0, linewidth=1.0)
    _apply_journal_style(
        ax,
        "Cumulative project cash flow at entered selling price",
        "Project time (year)",
        f"Cumulative cash flow ({ctx.currency})",
        grid_axis="both"
    )
    ax.yaxis.set_major_formatter(FuncFormatter(money_formatter))
    path = out_dir / f"{stem}_Cash_Flow.png"
    _save_figure(fig, path); paths.append(str(path))

    # Sensitivity ranking — combined NPV/MSP economic impact.
    if sensitivity is not None and not sensitivity.empty:
        p = sensitivity.sort_values("Impact Score", ascending=True).copy()
        fig, ax = plt.subplots(figsize=(10.2, max(6.0, len(p)*0.62 + 1.8)))
        y = np.arange(len(p))
        ax.barh(y, p["Impact Score"])
        ax.set_yticks(y)
        ax.set_yticklabels([_journal_label(v, 28) for v in p["Parameter"]])
        _apply_journal_style(
            ax,
            "Sensitivity ranking — combined economic impact",
            "Normalized impact score (software-defined)",
            None,
            grid_axis="x"
        )
        ax.set_xlim(0, 1.08)
        path = out_dir / f"{stem}_Sensitivity.png"
        _save_figure(fig, path); paths.append(str(path))

        # Separate tornado: low/high effect on MSP only.
        t = sensitivity.copy()
        active = (
            t["Low MSP Delta"].abs() + t["High MSP Delta"].abs()
        ) > 1e-10
        t = t[active].sort_values("Rank", ascending=False)
        if not t.empty:
            y = np.arange(len(t))
            fig, ax = plt.subplots(figsize=(10.5, max(6.0, len(t)*0.62 + 1.5)))
            ax.barh(y - 0.18, t["Low MSP Delta"], height=0.34, label="Low assumption")
            ax.barh(y + 0.18, t["High MSP Delta"], height=0.34, label="High assumption")
            ax.axvline(0, linewidth=1.0)
            ax.set_yticks(y)
            ax.set_yticklabels([_journal_label(v, 28) for v in t["Parameter"]])
            _apply_journal_style(
                ax,
                "Sensitivity tornado — change in minimum selling price",
                f"Change from baseline MSP ({ctx.currency}/product unit)",
                None,
                grid_axis="x"
            )
            ax.legend(frameon=False, fontsize=10.5, loc="lower right")
            path = out_dir / f"{stem}_Sensitivity_Tornado.png"
            _save_figure(fig, path); paths.append(str(path))

    # Scenario NPV.
    if scenarios is not None and not scenarios.empty:
        p = scenarios.copy()
        fig, ax = plt.subplots(figsize=(9.8, max(5.8, len(p)*0.65 + 2)))
        y = np.arange(len(p))
        ax.barh(y, p["NPV"])
        ax.axvline(0, linewidth=1.0)
        ax.set_yticks(y)
        ax.set_yticklabels([_journal_label(v, 24) for v in p["Scenario"]])
        _apply_journal_style(ax, "Scenario comparison", f"NPV ({ctx.currency})", None, "x")
        ax.xaxis.set_major_formatter(FuncFormatter(money_formatter))
        for yi, val in zip(y, p["NPV"]):
            ax.annotate(
                money_formatter(val), (val, yi), xytext=(5 if val >= 0 else -5, 0),
                textcoords="offset points", va="center",
                ha="left" if val >= 0 else "right", fontsize=9.5, fontweight="medium"
            )
        path = out_dir / f"{stem}_Scenarios.png"
        _save_figure(fig, path); paths.append(str(path))

    # Break-even margins.
    if break_even is not None and not break_even.empty:
        b = break_even[break_even["Status"].eq("Solved")].dropna(subset=["Change vs Base (%)"]).copy()
        if not b.empty:
            b = b.sort_values("Change vs Base (%)")
            fig, ax = plt.subplots(figsize=(10.2, max(5.8, len(b)*0.62 + 2)))
            y = np.arange(len(b))
            ax.barh(y, b["Change vs Base (%)"])
            ax.axvline(0, linewidth=1.0)
            ax.set_yticks(y)
            ax.set_yticklabels([_journal_label(v, 28) for v in b["Parameter"]])
            _apply_journal_style(ax, "Break-even margin from the base case", "Change from base assumption (%)", None, "x")
            path = out_dir / f"{stem}_Break_Even.png"
            _save_figure(fig, path); paths.append(str(path))

    # Scale MSP.
    if scale is not None and not scale.empty:
        s = scale.sort_values("Scale Factor")
        fig, ax = plt.subplots(figsize=(9.6, 5.9))
        ax.plot(s["Scale Factor"], s["MSP"], marker="o", markersize=5, linewidth=1.8)
        _apply_journal_style(
            ax, "Scale analysis — minimum selling price",
            "Plant scale factor", f"MSP ({ctx.currency}/product unit)", "both"
        )
        path = out_dir / f"{stem}_Scale.png"
        _save_figure(fig, path); paths.append(str(path))

    # Validation as percent deviation: avoids mixing unlike units on one axis.
    if validation is not None and not validation.empty:
        v = validation.dropna(subset=["Deviation (%)"]).copy()
        if not v.empty:
            v = v.sort_values("Deviation (%)")
            fig, ax = plt.subplots(figsize=(10.4, max(6.0, len(v)*0.60 + 2)))
            y = np.arange(len(v))
            ax.barh(y, v["Deviation (%)"])
            ax.axvline(0, linewidth=1.0)
            ax.axvline(1, linewidth=0.9, linestyle="--")
            ax.axvline(-1, linewidth=0.9, linestyle="--")
            ax.set_yticks(y)
            ax.set_yticklabels([_journal_label(x, 32) for x in v["Metric"]])
            _apply_journal_style(
                ax, "External validation — calculated deviation from published benchmark",
                "Deviation (%)", None, "x"
            )
            path = out_dir / f"{stem}_Validation.png"
            _save_figure(fig, path); paths.append(str(path))

    # Environmental emissions.
    if environmental is not None and not environmental.empty:
        names = ["Electricity emissions","Heat emissions","Feedstock emissions","Direct emissions"]
        e = environmental[environmental["Indicator"].isin(names)].copy()
        if not e.empty and pd.to_numeric(e["Value"], errors="coerce").fillna(0).sum() > 0:
            e["Value"] = pd.to_numeric(e["Value"], errors="coerce").fillna(0)
            e = e[e["Value"] > 0].sort_values("Value")
            if not e.empty:
                fig, ax = plt.subplots(figsize=(9.5, max(5.6, len(e)*0.7+2)))
                y = np.arange(len(e))
                ax.barh(y, e["Value"])
                ax.set_yticks(y)
                ax.set_yticklabels([_journal_label(x, 26) for x in e["Indicator"]])
                _apply_journal_style(
                    ax, "Environmental screening — annual emission contributions",
                    "Emissions (kg CO2e/year)", None, "x"
                )
                ax.xaxis.set_major_formatter(FuncFormatter(money_formatter))
                path = out_dir / f"{stem}_Environmental.png"
                _save_figure(fig, path); paths.append(str(path))

    # Uncertainty NPV distribution.
    if uncertainty_samples is not None and not uncertainty_samples.empty:
        u = uncertainty_samples
        fig, ax = plt.subplots(figsize=(9.8, 5.9))
        ax.hist(u["NPV"].dropna(), bins=36)
        for pct, style in [(10,"--"),(50,"-"),(90,":")]:
            val = float(np.percentile(u["NPV"].dropna(), pct))
            ax.axvline(val, linestyle=style, linewidth=1.5, label=f"P{pct}: {money_formatter(val)}")
        ax.axvline(0, linewidth=1.0)
        _apply_journal_style(
            ax, "Monte Carlo uncertainty — NPV distribution",
            f"NPV ({ctx.currency})", "Frequency", "both"
        )
        ax.xaxis.set_major_formatter(FuncFormatter(money_formatter))
        ax.legend(frameon=False, fontsize=10, loc="best")
        path = out_dir / f"{stem}_Uncertainty_NPV.png"
        _save_figure(fig, path); paths.append(str(path))

        fig, ax = plt.subplots(figsize=(9.8, 5.9))
        ax.hist(u["MSP"].dropna(), bins=36)
        for pct, style in [(10,"--"),(50,"-"),(90,":")]:
            val = float(np.percentile(u["MSP"].dropna(), pct))
            ax.axvline(val, linestyle=style, linewidth=1.5, label=f"P{pct}: {val:.3g}")
        ax.axvline(ctx.selling_price, linewidth=1.0, linestyle="-.", label="Entered selling price")
        _apply_journal_style(
            ax, "Monte Carlo uncertainty — MSP distribution",
            f"MSP ({ctx.currency}/product unit)", "Frequency", "both"
        )
        ax.legend(frameon=False, fontsize=10, loc="best")
        path = out_dir / f"{stem}_Uncertainty_MSP.png"
        _save_figure(fig, path); paths.append(str(path))

    # Uncertainty drivers.
    if uncertainty_drivers is not None and not uncertainty_drivers.empty:
        d = uncertainty_drivers.sort_values("Max |rho|", ascending=True)
        fig, ax = plt.subplots(figsize=(10.0, max(5.8, len(d)*0.62+2)))
        y = np.arange(len(d))
        ax.barh(y, d["Max |rho|"])
        ax.set_yticks(y)
        ax.set_yticklabels([_journal_label(x, 28) for x in d["Parameter"]])
        _apply_journal_style(
            ax, "Monte Carlo driver importance",
            "Maximum absolute Spearman correlation with NPV or MSP", None, "x"
        )
        ax.set_xlim(0, min(1.0, max(0.25, float(d["Max |rho|"].max())*1.12)))
        path = out_dir / f"{stem}_Uncertainty_Drivers.png"
        _save_figure(fig, path); paths.append(str(path))

    # R&D priority.
    if rd_priority is not None and not rd_priority.empty:
        r = rd_priority.sort_values("R&D Priority Score", ascending=True)
        fig, ax = plt.subplots(figsize=(10.3, max(6.0, len(r)*0.66+2)))
        y = np.arange(len(r))
        h = 0.34
        ax.barh(y-h/2, r["Risk Priority Score"], height=h, label="Economic risk priority")
        ax.barh(y+h/2, r["R&D Priority Score"], height=h, label="R&D priority")
        ax.set_yticks(y)
        ax.set_yticklabels([_journal_label(x, 28) for x in r["Parameter"]])
        _apply_journal_style(
            ax, "Risk and R&D priority ranking",
            "Normalized score (software-defined)", None, "x"
        )
        ax.set_xlim(0, 1.08)
        ax.legend(frameon=False, fontsize=10.5, loc="lower right")
        path = out_dir / f"{stem}_RnD_Priority.png"
        _save_figure(fig, path); paths.append(str(path))

    return paths


def format_results_workbook(path: Path):
    wb = load_workbook(path)
    navy = "17365D"
    light = "D9EAF7"
    thin = Side(style="thin", color="D9D9D9")

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 28

        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(color="FFFFFF", bold=True, size=11)
            cell.alignment = Alignment(wrap_text=True, vertical="center")

        for col in range(1, ws.max_column + 1):
            max_len = 0
            for row in range(1, min(ws.max_row, 120) + 1):
                val = ws.cell(row, col).value
                if val is not None:
                    max_len = max(max_len, len(str(val)))
            ws.column_dimensions[get_column_letter(col)].width = min(max(max_len + 2, 12), 42)

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="center")
                cell.border = Border(bottom=thin)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.00;[Red](#,##0.00);-'

        # Probability / correlation columns.
        for cell in ws[1]:
            h = str(cell.value or "").lower()
            if "probability" in h or "rho" in h or h == "irr":
                for r in range(2, ws.max_row + 1):
                    if isinstance(ws.cell(r, cell.column).value, (int, float)):
                        ws.cell(r, cell.column).number_format = '0.0%;[Red](0.0%);-'

        # Row-oriented KPI sheets: format fraction/probability values in column B.
        header_map = {str(c.value or "").strip().lower(): c.column for c in ws[1]}
        if "unit" in header_map and ("value" in header_map or "result" in header_map):
            value_col = header_map.get("value", header_map.get("result"))
            unit_col = header_map["unit"]
            for r in range(2, ws.max_row + 1):
                unit = str(ws.cell(r, unit_col).value or "").strip().lower()
                val_cell = ws.cell(r, value_col)
                if isinstance(val_cell.value, (int, float)) and unit in {"fraction","probability"}:
                    val_cell.number_format = '0.0%;[Red](0.0%);-'

    wb.save(path)


def write_results_excel(
    path: Path,
    ctx: ModelContext,
    baseline: Dict[str, object],
    sensitivity: pd.DataFrame,
    scenarios: pd.DataFrame,
    break_even: pd.DataFrame,
    scale: pd.DataFrame,
    environmental: pd.DataFrame,
    validation: pd.DataFrame,
    uncertainty_summary: Optional[pd.DataFrame] = None,
    uncertainty_samples: Optional[pd.DataFrame] = None,
    uncertainty_drivers: Optional[pd.DataFrame] = None,
    rd_priority: Optional[pd.DataFrame] = None,
    decision_summary: Optional[pd.DataFrame] = None,
):
    summary = pd.DataFrame([
        ["Case", ctx.case_name, "-"],
        ["Process category", ctx.process_category, "-"],
        ["Primary product", ctx.primary_product, "-"],
        ["Operating hours", ctx.hours, "h/year"],
        ["Project life", ctx.life, "years"],
        ["Purchased equipment / base capital", ctx.purchased_equipment, ctx.currency],
        ["Installed equipment / cost blocks", ctx.installed_equipment, ctx.currency],
        ["FCI", ctx.fci, ctx.currency],
        ["TCI", baseline["TCI"], ctx.currency],
        ["Annual product", baseline["Annual Product"], f"{ctx.product_unit.replace('/h','')}/year"],
        ["Annual OPEX", baseline["Annual OPEX"], f"{ctx.currency}/year"],
        ["Annual revenue", baseline["Annual Revenue"], f"{ctx.currency}/year"],
        ["NPV", baseline["NPV"], ctx.currency],
        ["IRR", baseline["IRR"], "fraction"],
        ["Payback", baseline["Payback"], "years"],
        ["MSP", baseline["MSP"], f"{ctx.currency}/product unit"],
        ["Entered selling price", ctx.selling_price, ctx.selling_price_unit],
        ["Benefit-cost ratio", baseline["BCR"], "ratio"],
        ["Finance mode", ctx.finance_mode, "-"],
    ], columns=["Metric", "Value", "Unit"])

    opex_df = pd.DataFrame(
        [{"OPEX Category": k, "Annual Cost": v} for k, v in baseline["OPEX Components"].items()]
    )
    warnings_df = pd.DataFrame({"Warning": ctx.warnings}) if ctx.warnings else pd.DataFrame({"Warning":["None"]})

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        if decision_summary is not None and not decision_summary.empty:
            decision_summary.to_excel(writer, sheet_name="Decision_Summary", index=False)
        ctx.equipment_df.to_excel(writer, sheet_name="Equipment_Costs", index=False)
        opex_df.to_excel(writer, sheet_name="OPEX_Breakdown", index=False)
        baseline["Cash Flow"].to_excel(writer, sheet_name="Cash_Flow", index=False)
        sensitivity.to_excel(writer, sheet_name="Sensitivity", index=False)
        scenarios.to_excel(writer, sheet_name="Scenarios", index=False)
        break_even.to_excel(writer, sheet_name="Break_Even", index=False)
        scale.to_excel(writer, sheet_name="Scale_Analysis", index=False)
        environmental.to_excel(writer, sheet_name="Environmental", index=False)
        if uncertainty_summary is not None:
            uncertainty_summary.to_excel(writer, sheet_name="Uncertainty_Summary", index=False)
        if uncertainty_drivers is not None:
            uncertainty_drivers.to_excel(writer, sheet_name="Uncertainty_Drivers", index=False)
        if rd_priority is not None:
            rd_priority.to_excel(writer, sheet_name="R&D_Priority", index=False)
        if uncertainty_samples is not None:
            uncertainty_samples.to_excel(writer, sheet_name="Uncertainty_Samples", index=False)
        validation.to_excel(writer, sheet_name="Validation_Comparison", index=False)
        warnings_df.to_excel(writer, sheet_name="Warnings", index=False)

    format_results_workbook(path)


def generate_verdict(
    ctx: ModelContext,
    baseline: Dict[str, object],
    scenarios: pd.DataFrame,
    sensitivity: pd.DataFrame,
    validation: pd.DataFrame,
    uncertainty_summary: Optional[pd.DataFrame] = None,
    rd_priority: Optional[pd.DataFrame] = None,
    decision_summary: Optional[pd.DataFrame] = None,
) -> str:
    robustness = None
    if decision_summary is not None and not decision_summary.empty:
        r = decision_summary[decision_summary["Decision Indicator"].eq("Economic robustness class")]
        if not r.empty:
            robustness = str(r.iloc[0]["Value"])

    top_sens = sensitivity.iloc[0]["Parameter"] if sensitivity is not None and not sensitivity.empty else "not evaluated"
    top_rd = rd_priority.iloc[0]["Parameter"] if rd_priority is not None and not rd_priority.empty else "not evaluated"

    prob_npv = np.nan
    if uncertainty_summary is not None and not uncertainty_summary.empty:
        prob_npv = _get_indicator(uncertainty_summary, "P(NPV > 0)")

    text = (
        f"The base case gives NPV = {baseline['NPV']:,.0f} {ctx.currency}, "
        f"IRR = {baseline['IRR']:.2%} and MSP = {baseline['MSP']:.4f} "
        f"{ctx.currency}/product unit. "
    )
    if robustness:
        text += f"The V0.6 software-defined robustness class is '{robustness}'. "
    if not np.isnan(prob_npv):
        text += f"Monte Carlo analysis gives P(NPV > 0) = {prob_npv:.1%}. "
    text += f"The highest deterministic sensitivity is {top_sens}. "
    if top_rd != "not evaluated":
        text += f"The highest R&D-priority parameter is {top_rd}. "

    if validation is not None and not validation.empty:
        m = validation[
            validation["Metric"].astype(str).str.lower().str.contains("production cost|msp", regex=True)
        ]
        if not m.empty and not np.isnan(m.iloc[0]["Deviation (%)"]):
            d = float(m.iloc[0]["Deviation (%)"])
            text += f"The primary external-validation deviation is {d:+.2f}%. "

    text += (
        "Uncertainty distributions and R&D-priority thresholds are user-defined decision-support assumptions; "
        "environmental results remain screening indicators rather than a full LCA."
    )
    return text


def write_word_report(
    path: Path,
    ctx: ModelContext,
    baseline: Dict[str, object],
    sensitivity: pd.DataFrame,
    scenarios: pd.DataFrame,
    break_even: pd.DataFrame,
    environmental: pd.DataFrame,
    validation: pd.DataFrame,
    uncertainty_summary: Optional[pd.DataFrame] = None,
    rd_priority: Optional[pd.DataFrame] = None,
    decision_summary: Optional[pd.DataFrame] = None,
):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.55)
    sec.bottom_margin = Inches(0.55)
    sec.left_margin = Inches(0.55)
    sec.right_margin = Inches(0.55)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(9.2)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.10

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Universal TEA V0.6 — Decision-Support Report")
    r.bold = True
    r.font.size = Pt(15)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(ctx.case_name)
    r.italic = True
    r.font.size = Pt(9)

    p = doc.add_paragraph()
    rr = p.add_run("Table 1. Key input assumptions")
    rr.bold = True; rr.font.size = Pt(11)
    inputs = [
        ["Process category", ctx.process_category, "-"],
        ["Operating hours", ctx.hours, "h/year"],
        ["Project life", ctx.life, "years"],
        ["Primary product", ctx.primary_product, "-"],
        ["Product rate", ctx.product_rate, ctx.product_unit],
        ["Selling price", ctx.selling_price, ctx.selling_price_unit],
        ["Target return", ctx.target_return, "fraction"],
        ["Tax rate", ctx.tax_rate, "fraction"],
        ["Finance mode", ctx.finance_mode, "-"],
        ["Installed capital", ctx.installed_equipment, ctx.currency],
    ]
    add_doc_table(doc, ["Input","Value","Unit"], inputs, 7.5)

    p = doc.add_paragraph()
    rr = p.add_run("Table 2. Main TEA and uncertainty outputs")
    rr.bold = True; rr.font.size = Pt(11)
    outputs = [
        ["TCI", baseline["TCI"], ctx.currency],
        ["Annual OPEX", baseline["Annual OPEX"], f"{ctx.currency}/year"],
        ["Annual revenue", baseline["Annual Revenue"], f"{ctx.currency}/year"],
        ["NPV", baseline["NPV"], ctx.currency],
        ["IRR", baseline["IRR"], "fraction"],
        ["Payback", baseline["Payback"], "years"],
        ["MSP", baseline["MSP"], f"{ctx.currency}/product unit"],
        ["Benefit-cost ratio", baseline["BCR"], "ratio"],
    ]
    if uncertainty_summary is not None and not uncertainty_summary.empty:
        for name in ["P10 NPV","P50 NPV","P90 NPV","P10 MSP","P50 MSP","P90 MSP",
                     "P(NPV > 0)","P(IRR >= target return)","P(MSP <= entered selling price)"]:
            row = uncertainty_summary[uncertainty_summary["Indicator"].eq(name)]
            if not row.empty:
                outputs.append([name, row.iloc[0]["Value"], row.iloc[0]["Unit"]])
    add_doc_table(doc, ["Output","Value","Unit"], outputs, 7.2)

    doc.add_page_break()
    p = doc.add_paragraph()
    rr = p.add_run("Table 3. Decision, risk and R&D-priority index")
    rr.bold = True; rr.font.size = Pt(11)

    rows = []
    if decision_summary is not None and not decision_summary.empty:
        chosen = [
            "Economic robustness class","Scenario viability rate","P(NPV > 0)",
            "Top deterministic sensitivity","Top uncertainty-risk parameter",
            "Top R&D priority","Closest break-even margin",
            "Carbon intensity","Carbon saving vs reference",
            "Primary external-validation deviation"
        ]
        for name in chosen:
            m = decision_summary[decision_summary["Decision Indicator"].eq(name)]
            if not m.empty:
                rrw = m.iloc[0]
                rows.append([name, rrw["Value"], rrw["Unit"], rrw["Module"]])
    if rd_priority is not None and not rd_priority.empty:
        for _, rrow in rd_priority.head(3).iterrows():
            rows.append([
                f"R&D rank {int(rrow['R&D Rank'])}",
                rrow["Parameter"],
                f"{rrow['R&D Priority Score']:.3f}",
                rrow["R&D Priority"]
            ])
    add_doc_table(doc, ["Index","Result","Unit / Score","Status / Module"], rows, 7.0)

    p = doc.add_paragraph()
    rr = p.add_run("Final verdict")
    rr.bold = True; rr.font.size = Pt(11)
    doc.add_paragraph(
        generate_verdict(
            ctx, baseline, scenarios, sensitivity, validation,
            uncertainty_summary, rd_priority, decision_summary
        )
    )

    if ctx.warnings:
        p = doc.add_paragraph()
        rr = p.add_run("Methodology warnings")
        rr.bold = True; rr.font.size = Pt(10)
        for warning in ctx.warnings:
            doc.add_paragraph(warning)

    doc.save(path)


def main():
    parser = argparse.ArgumentParser(
        description="Universal TEA Engine V0.6 — workbook-driven TEA, uncertainty and decision support"
    )
    parser.add_argument("input_workbook", help="Universal TEA V0.6 input .xlsx")
    parser.add_argument("--output-dir", default=".", help="Output directory")
    parser.add_argument("--no-report", action="store_true", help="Do not generate Word report")
    parser.add_argument("--no-figures", action="store_true", help="Do not generate PNG figures")
    args = parser.parse_args()

    input_path = Path(args.input_workbook)
    if not input_path.exists():
        raise SystemExit(f"Input workbook not found: {input_path}")

    out_dir = Path(args.output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    ctx, rd, decision_df, validation_targets = parse_model(str(input_path))
    stem = slugify(input_path.stem)

    baseline = evaluate_case(ctx)
    sensitivity = run_sensitivity(ctx, decision_df)
    scenarios = run_scenarios(ctx, decision_df)
    break_even = run_break_even(ctx, decision_df)
    scale = run_scale(ctx, decision_df)
    uncertainty_summary, uncertainty_samples, uncertainty_drivers = run_uncertainty(ctx, rd, baseline)
    rd_priority = build_rd_priority(ctx, baseline, sensitivity, rd, uncertainty_drivers)
    environmental = calculate_environmental(ctx, rd, baseline)
    validation = validation_comparison(ctx, validation_targets, baseline, sensitivity)
    decision_summary = build_decision_summary(
        ctx, baseline, sensitivity, scenarios, break_even,
        uncertainty_summary, rd_priority, environmental, validation, rd
    )

    results_path = out_dir / f"{stem}_TEA_Results_V0_6.xlsx"
    report_path = out_dir / f"{stem}_TEA_Report_V0_6.docx"

    write_results_excel(
        results_path, ctx, baseline, sensitivity, scenarios, break_even, scale,
        environmental, validation, uncertainty_summary, uncertainty_samples,
        uncertainty_drivers, rd_priority, decision_summary
    )

    if not args.no_report:
        write_word_report(
            report_path, ctx, baseline, sensitivity, scenarios, break_even,
            environmental, validation, uncertainty_summary, rd_priority,
            decision_summary
        )

    figure_paths = []
    if not args.no_figures:
        figure_paths = save_charts(
            ctx, out_dir, stem, baseline, sensitivity, scenarios, validation,
            break_even=break_even, scale=scale, environmental=environmental,
            uncertainty_summary=uncertainty_summary,
            uncertainty_samples=uncertainty_samples,
            uncertainty_drivers=uncertainty_drivers,
            rd_priority=rd_priority
        )

    print("=" * 78)
    print("UNIVERSAL TEA ENGINE V0.6")
    print("=" * 78)
    print(f"Case                    : {ctx.case_name}")
    print(f"Finance mode            : {ctx.finance_mode}")
    print(f"TCI                     : {baseline['TCI']:,.2f} {ctx.currency}")
    print(f"Annual OPEX             : {baseline['Annual OPEX']:,.2f} {ctx.currency}/year")
    print(f"Annual revenue          : {baseline['Annual Revenue']:,.2f} {ctx.currency}/year")
    print(f"NPV                     : {baseline['NPV']:,.2f} {ctx.currency}")
    print(f"IRR                     : {baseline['IRR']:.4%}" if not np.isnan(baseline["IRR"]) else "IRR                     : n/a")
    print(f"Payback                 : {baseline['Payback']:.3f} years" if not np.isnan(baseline["Payback"]) else "Payback                 : not reached")
    print(f"MSP                     : {baseline['MSP']:.6f} {ctx.currency}/product unit")

    if not uncertainty_summary.empty:
        p_npv = _get_indicator(uncertainty_summary, "P(NPV > 0)")
        p_msp = _get_indicator(uncertainty_summary, "P(MSP <= entered selling price)")
        print(f"P(NPV > 0)              : {p_npv:.2%}")
        print(f"P(MSP <= selling price) : {p_msp:.2%}")

    if not rd_priority.empty:
        top = rd_priority.iloc[0]
        print(f"Top R&D priority        : {top['Parameter']} (score {top['R&D Priority Score']:.3f})")

    if not validation.empty:
        print("\nValidation comparison:")
        print(validation[["Metric","Published","Calculated","Deviation (%)","Status"]].to_string(index=False))

    if ctx.warnings:
        print("\nWarnings:")
        for w in ctx.warnings:
            print(" -", w)

    print("\nOutputs:")
    print(f" - {results_path}")
    if not args.no_report:
        print(f" - {report_path}")
    for p in figure_paths:
        print(f" - {p}")


if __name__ == "__main__":
    main()
